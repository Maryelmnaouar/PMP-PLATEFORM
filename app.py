# -*- coding: utf-8 -*-
"""
ASSETPILOT — Maintenance Autonome
Flask + SQLite. Configuration par variables d'environnement (voir README.md et .env.example).
Local : python app.py   ·   Production : gunicorn app:app
"""
import os
import json
import random
import sqlite3
import datetime as dt
from functools import wraps

from flask import (Flask, render_template, request, redirect, url_for, flash,
                   session, abort, g, jsonify, send_from_directory)
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename
from werkzeug.middleware.proxy_fix import ProxyFix

import sap_integration as sap
import pg_store


def _load_dotenv(path):
    """Charge un fichier .env local s'il existe (les variables déjà définies sont prioritaires)."""
    if not os.path.exists(path):
        return
    for line in open(path, encoding="utf-8"):
        line = line.split(" #")[0].strip()
        if line and not line.startswith("#") and "=" in line:
            k, v = line.split("=", 1)
            os.environ.setdefault(k.strip(), v.strip())


_load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))

BASE = os.path.dirname(os.path.abspath(__file__))
DATA_DIR = os.path.abspath(os.environ.get("DATA_DIR", BASE))      # sur Render : le disque persistant
DB = os.path.join(DATA_DIR, "assetpilot.db")
UPLOADS = os.path.join(DATA_DIR, "uploads")
SCHEMA_VERSION = "6.3"
MIGRATABLE_VERSIONS = ("6.0", "6.1", "6.2")  # mises à jour sans perte de données

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)   # derrière le proxy HTTPS de Render
_secret = os.environ.get("SECRET_KEY")
if not _secret:
    import secrets as _secrets
    _secret = _secrets.token_hex(32)
    print("ATTENTION : SECRET_KEY non définie — clé temporaire (les sessions seront perdues au redémarrage).")
app.secret_key = _secret
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
app.config["SESSION_COOKIE_SECURE"] = os.environ.get("SESSION_COOKIE_SECURE", "0") == "1"
app.config["PERMANENT_SESSION_LIFETIME"] = dt.timedelta(hours=12)

# Compte administrateur créé au premier démarrage (base vide) — jamais de mot de passe dans le code
ADMIN_USER = os.environ.get("ADMIN_USERNAME", "admin")
ADMIN_PASSWORD = os.environ.get("ADMIN_PASSWORD", "")
ADMIN_NAME = os.environ.get("ADMIN_NAME", "Administrateur")
ADMIN_EMAIL = os.environ.get("ADMIN_EMAIL", "")

# ---------------------------------------------------------------------------
# Référentiels
# ---------------------------------------------------------------------------
ROLES = [
    "Administrateur", "Opérateur", "Technicien",
    "Chef d'équipe Production", "Chef d'équipe Maintenance", "Leader SGA",
    "Plant Manager", "Groupe OE",
    "Responsable magasin", "Responsable maintenance", "Expert ligne",
]
FIELD_ROLES = ("Opérateur", "Technicien")
TEAM_ROLES = ("Chef d'équipe Production", "Chef d'équipe Maintenance", "Leader SGA")
CHEF_ROLES = ("Chef d'équipe Production", "Chef d'équipe Maintenance")
CHEF_OF = {"Opérateur": "Chef d'équipe Production", "Technicien": "Chef d'équipe Maintenance"}
PLANT_ROLES = ("Plant Manager", "Responsable magasin", "Responsable maintenance", "Expert ligne")
GLOBAL_ROLES = ("Administrateur", "Groupe OE")

TAG_TYPES = [
    {"key": "AM", "label": "AM", "color_name": "Bleu", "color": "#1687e8", "head": "#0d1f4f", "text": "#fff",
     "hint": "Anomalie traitable par l'opérateur"},
    {"key": "RCM", "label": "RCM", "color_name": "Rouge", "color": "#d71920", "head": "#d71920", "text": "#fff",
     "hint": "Nécessite l'intervention de la maintenance"},
    {"key": "SAFETY", "label": "Safety", "color_name": "Jaune", "color": "#f4c400", "head": "#f4c400", "text": "#1d1d1d",
     "hint": "Risque sécurité / HSE"},
]
TAG_TYPE = {t["key"]: t for t in TAG_TYPES}

STATUSES = ["Nouveau", "En cours", "Traité", "Contre-mesure"]
STATUS_INFO = {
    "Nouveau": {"color": "#d71920", "text": "#fff", "board": "Anomalies découvertes"},
    "En cours": {"color": "#f08a00", "text": "#fff", "board": "Anomalies en cours de traitement"},
    "Traité": {"color": "#58c48a", "text": "#0b3d22", "board": "Anomalies traitées"},
    "Contre-mesure": {"color": "#0b5e34", "text": "#fff", "board": "Contre-mesure (action préventive)"},
}

ANOMALY_TYPES = [
    "Défauts mineurs",
    "Absence des conditions de base",
    "Zones difficiles d'accès (HTR – Hard to Reach)",
    "Sources de contamination",
    "Défauts qualité",
    "Équipements / éléments inutiles",
    "Conditions dangereuses / situations à risque",
]
SOC_HTR_TYPES = ("Sources de contamination", "Zones difficiles d'accès (HTR – Hard to Reach)")

FREQUENCIES = ["Quotidienne", "Hebdomadaire", "Mensuelle", "Trimestrielle", "Semestrielle", "Annuelle"]
FREQ_STEP = {"Quotidienne": ("days", 1), "Hebdomadaire": ("days", 7), "Mensuelle": ("months", 1),
             "Trimestrielle": ("months", 3), "Semestrielle": ("months", 6), "Annuelle": ("months", 12)}
PLAN_PERIODS = {"semaine": ("Semaine", ["Quotidienne", "Hebdomadaire"]),
                "mois": ("Mois", ["Mensuelle", "Trimestrielle"]),
                "annee": ("Année", ["Semestrielle", "Annuelle"])}
CIL_TYPES = [("C", "Cleaning", "Nettoyage", "#1687e8"), ("I", "Inspection", "Inspection", "#f08a00"),
             ("L", "Lubrification", "Lubrification", "#17a45a"), ("S", "Serrage", "Serrage (Tightening)", "#6b4fd8")]
CIL = {c[0]: {"key": c[0], "label": c[1], "fr": c[2], "color": c[3]} for c in CIL_TYPES}
MACHINE_STATES = ["Machine à l'arrêt", "Machine en marche"]
DOC_TYPES = [("SOP", "SOP"), ("OPL", "OPL"), ("CHECK", "Checklist"), ("TECH", "Fiche technique")]
TASK_OPEN = ("À réaliser", "En cours")
TASK_DONE = ("Réalisée", "Confirmée")

PERMS = [
    ("tags_create", "Lancer un TAG"),
    ("tags_move", "Changer le statut des TAGs de son périmètre"),
    ("tags_all", "Agir sur tous les TAGs"),
    ("tag_validate", "Valider les TAGs RCM / Safety (N° d'avis SAP)"),
    ("countermeasure", "Remplir le formulaire de contre-mesure"),
    ("tasks_exec", "Exécuter ses tâches / checklists"),
    ("confirm", "Confirmer les travaux de l'équipe"),
    ("checklists_manage", "Gérer les checklists CILT, les contre-mesures et la distribution des tâches"),
    ("docs_manage", "Gérer la documentation"),
    ("sga_board_edit", "Modifier le tableau SGA"),
    ("dashboard_global", "Tableau de bord global (périmètre)"),
    ("audit_self", "Réaliser un self audit"),
    ("audit_pm", "Réaliser un audit Plant Manager"),
    ("audit_oe", "Réaliser un audit Groupe OE"),
    ("pdr_view", "Voir la gestion PDR"),
    ("pdr_request", "Créer une demande PDR"),
    ("pdr_validate", "Valider une demande PDR"),
    ("pdr_order", "Commander / réceptionner les PDR (magasin)"),
    ("settings", "Paramètres & utilisateurs"),
]
PERM_KEYS = [p[0] for p in PERMS]
ROLE_DEFAULTS = {
    "Administrateur": PERM_KEYS,
    "Opérateur": ["tags_create", "tags_move", "tasks_exec"],
    "Technicien": ["tags_create", "tags_move", "tasks_exec"],
    "Chef d'équipe Production": ["tags_create", "tags_move", "tasks_exec", "confirm", "tag_validate"],
    "Chef d'équipe Maintenance": ["tags_create", "tags_move", "tasks_exec", "confirm"],
    "Leader SGA": ["tags_create", "tags_move", "tasks_exec", "audit_self", "sga_board_edit"],
    "Plant Manager": ["tags_create", "dashboard_global", "audit_pm"],
    "Groupe OE": ["tags_create", "dashboard_global", "audit_oe"],
    "Responsable magasin": ["tags_create", "pdr_view", "pdr_order"],
    "Responsable maintenance": ["tags_create", "tags_move", "dashboard_global", "pdr_view", "pdr_request", "pdr_validate"],
    "Expert ligne": ["tags_create", "tags_move", "pdr_view", "pdr_request"],
}

AUDIT_LEVELS = [
    {"key": "Self", "label": "Self audit", "perm": "audit_self", "color": "#1687e8",
     "who": "Chaque SGA s'auto-évalue", "scope": "Sa propre SGA"},
    {"key": "Plant Manager", "label": "Audit Plant Manager", "perm": "audit_pm", "color": "#f08a00",
     "who": "Le Plant Manager audite les SGA de son usine", "scope": "Toutes les SGA de l'usine"},
    {"key": "Groupe OE", "label": "Audit Groupe OE", "perm": "audit_oe", "color": "#d71920",
     "who": "L'équipe Groupe OE audite toutes les usines", "scope": "Global"},
]
AUDIT_THRESHOLD = 80
AUDIT_GRID = {
    1: [("Nettoyage initial réalisé sur tous les équipements de la SGA", None),
        ("Anomalies détectées et étiquetées (TAGs émis)", "tags_count"),
        ("Taux de traitement des étiquettes ≥ 80 %", "tag_rate"),
        ("Sources de contamination et zones difficiles d'accès identifiées", "soc_htr"),
        ("Standards provisoires de nettoyage et d'inspection rédigés", "sop_count"),
        ("Tableau SGA tenu à jour", "board_update"),
        ("Membres formés à la sécurité et aux règles d'étiquetage", None)],
    2: [("Sources de contamination éliminées ou maîtrisées", "soc_htr"),
        ("Zones difficiles d'accès améliorées", None),
        ("Temps de nettoyage / inspection réduit", "ci_time"),
        ("Contre-mesures en place pour les anomalies récurrentes", "cm_count"),
        ("Taux de traitement des étiquettes ≥ 80 %", "tag_rate"),
        ("Tableau SGA tenu à jour", "board_update"),
        ("Kaizens réalisés et documentés", None)],
    3: [("Standards CIL (nettoyage, inspection, lubrification) établis", "sop_count"),
        ("Checklists exécutées selon la fréquence prévue", "realization"),
        ("OPL créées et diffusées", "opl_count"),
        ("Points de lubrification identifiés et repérés", None),
        ("Transfert d'activités RCM → AM engagé", "transfer"),
        ("Respect des temps standards", "ci_time"),
        ("Tableau SGA tenu à jour", "board_update")],
}
AUDIT_GRID_DEFAULT = [
    ("Objectifs de l'étape atteints", None),
    ("Standards à jour et respectés", "sop_count"),
    ("Checklists exécutées selon la fréquence prévue", "realization"),
    ("Taux de traitement des étiquettes ≥ 80 %", "tag_rate"),
    ("Tableau SGA tenu à jour", "board_update"),
    ("Implication de toute l'équipe SGA", None),
]

PDR_STATUSES = ["Demandée", "Validée", "Commandée", "Réceptionnée", "Refusée"]
PDR_COLORS = {"Demandée": "#6b7480", "Validée": "#f08a00", "Commandée": "#1687e8",
              "Réceptionnée": "#17a45a", "Refusée": "#d71920"}

# ---------------------------------------------------------------------------
# Base de données
# ---------------------------------------------------------------------------
SCHEMA = """
CREATE TABLE IF NOT EXISTS meta(key TEXT PRIMARY KEY, value TEXT);
CREATE TABLE IF NOT EXISTS plants(
    id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT UNIQUE NOT NULL, name TEXT NOT NULL,
    city TEXT, status TEXT DEFAULT 'Actif');
CREATE TABLE IF NOT EXISTS lines(
    id INTEGER PRIMARY KEY AUTOINCREMENT, plant_id INTEGER NOT NULL REFERENCES plants(id) ON DELETE CASCADE,
    code TEXT, name TEXT NOT NULL);
CREATE TABLE IF NOT EXISTS sga(
    id INTEGER PRIMARY KEY AUTOINCREMENT, line_id INTEGER NOT NULL REFERENCES lines(id) ON DELETE CASCADE,
    name TEXT NOT NULL, description TEXT, am_step INTEGER DEFAULT 1, created_at TEXT);
CREATE TABLE IF NOT EXISTS equipment(
    id INTEGER PRIMARY KEY AUTOINCREMENT, line_id INTEGER NOT NULL REFERENCES lines(id) ON DELETE CASCADE,
    sga_id INTEGER REFERENCES sga(id) ON DELETE SET NULL, reference TEXT, name TEXT NOT NULL,
    position INTEGER DEFAULT 0, description TEXT);
CREATE TABLE IF NOT EXISTS users(
    id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, username TEXT UNIQUE NOT NULL,
    email TEXT, role TEXT NOT NULL, password_hash TEXT NOT NULL, plant_id INTEGER REFERENCES plants(id),
    permissions TEXT, active INTEGER DEFAULT 1, must_change_password INTEGER DEFAULT 0,
    supervisor_id INTEGER REFERENCES users(id) ON DELETE SET NULL);
CREATE TABLE IF NOT EXISTS user_sga(
    user_id INTEGER REFERENCES users(id) ON DELETE CASCADE, sga_id INTEGER REFERENCES sga(id) ON DELETE CASCADE,
    PRIMARY KEY(user_id, sga_id));
CREATE TABLE IF NOT EXISTS tags(
    id INTEGER PRIMARY KEY AUTOINCREMENT, tag_code TEXT UNIQUE, type TEXT NOT NULL, anomaly_type TEXT,
    equipment_id INTEGER NOT NULL REFERENCES equipment(id), sga_id INTEGER, line_id INTEGER,
    organe TEXT, description TEXT NOT NULL, photo TEXT,
    observed_by INTEGER REFERENCES users(id), observed_at TEXT, created_at TEXT, ot_number TEXT,
    treated_by INTEGER REFERENCES users(id), treatment_start TEXT, resolved_at TEXT,
    status TEXT DEFAULT 'Nouveau', countermeasure TEXT,
    confirmed_by INTEGER, confirmed_at TEXT, updated_at TEXT);
CREATE TABLE IF NOT EXISTS tag_history(
    id INTEGER PRIMARY KEY AUTOINCREMENT, tag_id INTEGER REFERENCES tags(id) ON DELETE CASCADE,
    user_id INTEGER, from_status TEXT, to_status TEXT, comment TEXT, at TEXT);
CREATE TABLE IF NOT EXISTS checklists(
    id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT, name TEXT NOT NULL,
    equipment_id INTEGER NOT NULL REFERENCES equipment(id) ON DELETE CASCADE, organe TEXT,
    frequency TEXT DEFAULT 'Quotidienne', duration_min INTEGER DEFAULT 10, machine_state TEXT,
    people INTEGER DEFAULT 1, sop_url TEXT, activity_type TEXT DEFAULT 'AM',
    assigned_to INTEGER REFERENCES users(id), origin_tag_id INTEGER REFERENCES tags(id),
    countermeasure TEXT, transferred INTEGER DEFAULT 0, transfer_date TEXT,
    active INTEGER DEFAULT 1, created_at TEXT);
CREATE TABLE IF NOT EXISTS checklist_items(
    id INTEGER PRIMARY KEY AUTOINCREMENT, checklist_id INTEGER REFERENCES checklists(id) ON DELETE CASCADE,
    position INTEGER, text TEXT NOT NULL, organe TEXT, cil_type TEXT, tooling TEXT, duration_min INTEGER,
    frequency TEXT, last_due TEXT, last_assignee INTEGER);
CREATE TABLE IF NOT EXISTS tasks(
    id INTEGER PRIMARY KEY AUTOINCREMENT, checklist_id INTEGER REFERENCES checklists(id) ON DELETE CASCADE,
    assigned_to INTEGER REFERENCES users(id), due_date TEXT, status TEXT DEFAULT 'À réaliser',
    started_at TEXT, completed_at TEXT, duration_min INTEGER, confirmed_by INTEGER, confirmed_at TEXT,
    reject_comment TEXT);
CREATE TABLE IF NOT EXISTS task_items(
    id INTEGER PRIMARY KEY AUTOINCREMENT, task_id INTEGER REFERENCES tasks(id) ON DELETE CASCADE,
    position INTEGER, text TEXT, ok INTEGER, validated_at TEXT, tag_id INTEGER, checklist_item_id INTEGER,
    organe TEXT, cil_type TEXT, tooling TEXT, duration_min INTEGER, frequency TEXT);
CREATE TABLE IF NOT EXISTS countermeasures(
    id INTEGER PRIMARY KEY AUTOINCREMENT, tag_id INTEGER, equipment_id INTEGER REFERENCES equipment(id) ON DELETE CASCADE,
    organe TEXT, countermeasure TEXT NOT NULL, tasks TEXT, cil_type TEXT, tooling TEXT, duration_min INTEGER,
    frequency TEXT, machine_state TEXT, people INTEGER DEFAULT 1, sop_url TEXT, activity_type TEXT DEFAULT 'AM',
    status TEXT DEFAULT 'À intégrer', checklist_id INTEGER, integrated_at TEXT, integrated_by INTEGER,
    created_by INTEGER, created_at TEXT);
CREATE TABLE IF NOT EXISTS documents(
    id INTEGER PRIMARY KEY AUTOINCREMENT, equipment_id INTEGER REFERENCES equipment(id) ON DELETE CASCADE,
    title TEXT NOT NULL, doc_type TEXT, version TEXT, owner TEXT, url TEXT, updated_at TEXT);
CREATE TABLE IF NOT EXISTS sga_board(
    sga_id INTEGER PRIMARY KEY REFERENCES sga(id) ON DELETE CASCADE,
    charte TEXT, objectifs TEXT, centrage TEXT, score_5s INTEGER, plan_5s TEXT, kaizen TEXT,
    mtbf TEXT, mttr TEXT, pertes_cadence TEXT, pannes TEXT, arrets_mineurs TEXT, preparatifs TEXT,
    updated_at TEXT, updated_by INTEGER);
CREATE TABLE IF NOT EXISTS audits(
    id INTEGER PRIMARY KEY AUTOINCREMENT, level TEXT, sga_id INTEGER REFERENCES sga(id) ON DELETE CASCADE,
    am_step INTEGER, auditor_id INTEGER, audit_date TEXT, score REAL, comment TEXT);
CREATE TABLE IF NOT EXISTS audit_items(
    id INTEGER PRIMARY KEY AUTOINCREMENT, audit_id INTEGER REFERENCES audits(id) ON DELETE CASCADE,
    criterion TEXT, note INTEGER, proof TEXT);
CREATE TABLE IF NOT EXISTS spare_parts(
    id INTEGER PRIMARY KEY AUTOINCREMENT, reference TEXT UNIQUE, designation TEXT, stock REAL DEFAULT 0,
    min_stock REAL DEFAULT 0, unit TEXT DEFAULT 'pcs', location TEXT);
CREATE TABLE IF NOT EXISTS pdr_requests(
    id INTEGER PRIMARY KEY AUTOINCREMENT, code TEXT, part_id INTEGER REFERENCES spare_parts(id),
    designation TEXT, qty REAL, equipment_id INTEGER, tag_id INTEGER, requester_id INTEGER,
    status TEXT DEFAULT 'Demandée', comment TEXT, created_at TEXT, validated_by INTEGER, validated_at TEXT,
    ordered_at TEXT, expected_date TEXT, received_at TEXT);
"""


def connect():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def db():
    if "db" not in g:
        g.db = connect()
    return g.db


@app.teardown_appcontext
def close_db(exc):
    conn = g.pop("db", None)
    if conn is not None:
        conn.close()


def q(sql, args=(), one=False):
    cur = db().execute(sql, args)
    rows = cur.fetchall()
    return (rows[0] if rows else None) if one else rows


def ex(sql, args=()):
    cur = db().execute(sql, args)
    db().commit()
    return cur.lastrowid


def now():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M")


def today():
    return dt.date.today()


def fmt_date(value, with_time=False):
    if not value:
        return "—"
    try:
        if len(value) > 10:
            d = dt.datetime.strptime(value[:16], "%Y-%m-%d %H:%M")
            return d.strftime("%d/%m/%Y %H:%M") if with_time else d.strftime("%d/%m/%Y")
        return dt.datetime.strptime(value, "%Y-%m-%d").strftime("%d/%m/%Y")
    except ValueError:
        return value


app.jinja_env.filters["d"] = fmt_date
app.jinja_env.filters["dt"] = lambda v: fmt_date(v, True)
app.jinja_env.filters["short"] = lambda n: (n or "").replace(" - CSD PET 4", "").replace("- CSD PET 4", "").strip()


# ---------------------------------------------------------------------------
# Initialisation de la base et données de départ
# ---------------------------------------------------------------------------
# Données de base réelles : usine Cobomi, ligne MOR_CA_CSD_PET4 et ses 26 équipements (références SAP).
# Désactivables avec la variable d'environnement SEED_BASE_DATA=0 (base totalement vide, à part l'administrateur).
BASE_PLANT = ("COB", "Cobomi", "Casablanca")
BASE_LINE = ("MOR_CA_CSD_PET4", "MOR_CA_CSD_PET4")
EQUIPMENT_COBOMI = [
    ("5001126", "ALIMENTATION DE PREFORMES"), ("5001127", "INSPECTRICE PREFORMES"),
    ("5001128", "RINCEUSE PREFORME"), ("5001129", "SOUFFLEUSE"),
    ("5001130", "INSPECTRICE BTLLES SOUFFLEES"), ("5001131", "SOUTIREUSE"),
    ("5001132", "VISSEUSE"), ("5001133", "MIXEUR"), ("5001134", "COBRIX"),
    ("5001135", "TRANSPORTEUR DE BOUCHONS"), ("5001136", "SECHEUR"),
    ("5001137", "INSPECTRICE NIVEAU"), ("5001138", "TRANSPORTEUR DE RECIPIENTS"),
    ("5001139", "ACCUMULATEUR"), ("5001140", "ETIQUETEUSE"), ("5001141", "INSPECTRICE ETIQUETTE"),
    ("5001142", "DATEUSE"), ("5001143", "FARDELEUSE"), ("5001144", "TRANSPORTEUR D'EMBALLAGES"),
    ("5001145", "ROBOT BOX"), ("5001146", "TRANSPORTEUR DE PALETTES"), ("5001147", "PALETTISEUR"),
    ("5001148", "STRECHEUSE"), ("5001149", "ETIQUETAGE PALETTE - CSD PET 4 LOGOPACK"),
    ("5001150", "IMPRIMANTE SICPA"), ("5001151", "STAT DE LUBRIF CONV A CHAINE"),
]


def _cols(conn, table):
    return {r[1] for r in conn.execute("PRAGMA table_info(%s)" % table).fetchall()}


def migrate_62(conn, upgrade=False):
    """6.1 → 6.2 : checklists CILT par tâche, contre-mesures séparées, mot de passe à changer, validation RCM/Safety."""
    conn.executescript(SCHEMA)  # crée les nouvelles tables (countermeasures) si absentes
    adds = {"users": [("must_change_password", "INTEGER DEFAULT 0")],
            "checklist_items": [("organe", "TEXT"), ("cil_type", "TEXT"), ("tooling", "TEXT"), ("duration_min", "INTEGER"),
                                ("frequency", "TEXT"), ("last_due", "TEXT"), ("last_assignee", "INTEGER")],
            "task_items": [("checklist_item_id", "INTEGER"), ("organe", "TEXT"), ("cil_type", "TEXT"), ("tooling", "TEXT"),
                           ("duration_min", "INTEGER"), ("frequency", "TEXT")]}
    fresh_items = "frequency" not in _cols(conn, "checklist_items")
    for table, cols in adds.items():
        have = _cols(conn, table)
        for name, typ in cols:
            if name not in have:
                conn.execute("ALTER TABLE %s ADD COLUMN %s %s" % (table, name, typ))
    if fresh_items:
        # reprise des réglages de la checklist sur chacune de ses tâches
        for cl in conn.execute("SELECT * FROM checklists").fetchall():
            items = conn.execute("SELECT id FROM checklist_items WHERE checklist_id=?", (cl["id"],)).fetchall()
            per = max(1, round((cl["duration_min"] or 10) / max(1, len(items))))
            last = conn.execute("SELECT MAX(due_date) FROM tasks WHERE checklist_id=?", (cl["id"],)).fetchone()[0]
            for it in items:
                conn.execute("UPDATE checklist_items SET organe=?, frequency=?, duration_min=?, last_due=?, last_assignee=? WHERE id=?",
                             (cl["organe"], cl["frequency"] or "Hebdomadaire", per, last, cl["assigned_to"], it["id"]))
        # contre-mesures déjà enregistrées : reprises dans la nouvelle liste (déjà intégrées)
        for cl in conn.execute("SELECT * FROM checklists WHERE origin_tag_id IS NOT NULL").fetchall():
            if not conn.execute("SELECT 1 FROM countermeasures WHERE tag_id=?", (cl["origin_tag_id"],)).fetchone():
                pts = "\n".join(r[0] for r in conn.execute("SELECT text FROM checklist_items WHERE checklist_id=? ORDER BY position", (cl["id"],)))
                conn.execute("""INSERT INTO countermeasures(tag_id,equipment_id,organe,countermeasure,tasks,duration_min,frequency,
                                machine_state,people,sop_url,activity_type,status,checklist_id,integrated_at,created_at)
                                VALUES(?,?,?,?,?,?,?,?,?,?,?, 'Intégrée', ?, ?, ?)""",
                             (cl["origin_tag_id"], cl["equipment_id"], cl["organe"], cl["countermeasure"] or cl["name"], pts,
                              cl["duration_min"], cl["frequency"], cl["machine_state"], cl["people"], cl["sop_url"],
                              cl["activity_type"], cl["id"], cl["created_at"], cl["created_at"]))
    # nouveau droit : validation des TAGs RCM / Safety pour les chefs d'équipe production (lors de la mise à jour)
    for u in [] if not upgrade else conn.execute("SELECT id, permissions FROM users WHERE role=?", ("Chef d'équipe Production",)).fetchall():
        perms = json.loads(u["permissions"] or "[]")
        if "tag_validate" not in perms:
            perms.append("tag_validate")
            conn.execute("UPDATE users SET permissions=? WHERE id=?", (json.dumps(perms), u["id"]))
    conn.commit()


def migrate_63(conn, upgrade=False):
    """6.2 → 6.3 : chaque opérateur / technicien est rattaché directement à son chef d'équipe."""
    if "supervisor_id" not in _cols(conn, "users"):
        conn.execute("ALTER TABLE users ADD COLUMN supervisor_id INTEGER")
        upgrade = True
    if upgrade:
        # proposition de départ : le chef d'équipe unique partageant une SGA avec l'opérateur / le technicien
        for role, chef_role in (("Opérateur", "Chef d'équipe Production"), ("Technicien", "Chef d'équipe Maintenance")):
            for u in conn.execute("SELECT id FROM users WHERE role=? AND supervisor_id IS NULL", (role,)).fetchall():
                chefs = conn.execute("""SELECT DISTINCT c.id FROM users c JOIN user_sga a ON a.user_id=c.id JOIN user_sga b ON b.sga_id=a.sga_id
                                        WHERE b.user_id=? AND c.role=? AND c.active=1""", (u["id"], chef_role)).fetchall()
                if len(chefs) == 1:
                    conn.execute("UPDATE users SET supervisor_id=? WHERE id=?", (chefs[0]["id"], u["id"]))
    conn.commit()


def migrate_concours(conn):
    """Concours de nom de l'application (tables créées si absentes, sans toucher au reste)."""
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS contest_names (
        id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL, reason TEXT, user_id INTEGER REFERENCES users(id),
        hidden INTEGER DEFAULT 0, created_at TEXT NOT NULL);
    CREATE TABLE IF NOT EXISTS contest_votes (
        user_id INTEGER PRIMARY KEY REFERENCES users(id), name_id INTEGER NOT NULL REFERENCES contest_names(id),
        voted_at TEXT NOT NULL);
    INSERT OR IGNORE INTO meta(key, value) VALUES('contest_phase', 'propositions');
    """)
    conn.commit()


def init_db():
    os.makedirs(os.path.dirname(DB), exist_ok=True)
    if os.path.exists(DB):
        conn = connect()
        try:
            row = conn.execute("SELECT value FROM meta WHERE key='schema'").fetchone()
        except sqlite3.Error:
            row = None
        if row and row[0] == SCHEMA_VERSION:
            sap.migrate(conn)
            migrate_62(conn)
            migrate_63(conn)
            migrate_concours(conn)
            conn.close()
            return
        if row and row[0] in MIGRATABLE_VERSIONS:
            sap.migrate(conn)
            migrate_62(conn, upgrade=row[0] in ("6.0", "6.1"))
            migrate_63(conn, upgrade=True)
            migrate_concours(conn)
            conn.execute("UPDATE meta SET value=? WHERE key='schema'", (SCHEMA_VERSION,))
            conn.commit()
            conn.close()
            print("Base mise à jour en version", SCHEMA_VERSION, "(données conservées)")
            return
        conn.close()
        raise SystemExit("La base %s n'est pas une base AssetPilot %s. Elle n'a pas été modifiée : "
                         "déplacez-la ou indiquez un autre DATA_DIR." % (DB, SCHEMA_VERSION))
    conn = connect()
    conn.executescript(SCHEMA)
    conn.execute("INSERT OR REPLACE INTO meta(key,value) VALUES('schema',?)", (SCHEMA_VERSION,))
    seed(conn)
    sap.migrate(conn)
    migrate_62(conn)
    migrate_63(conn)
    migrate_concours(conn)
    conn.commit()
    conn.close()


def seed(conn):
    """Première installation : compte administrateur + données de base (sans aucune donnée d'exemple)."""
    password = ADMIN_PASSWORD
    if not password:
        import secrets as _secrets
        password = _secrets.token_urlsafe(12)
        print("=" * 70)
        print(" Compte administrateur créé : %s" % ADMIN_USER)
        print(" Mot de passe provisoire    : %s" % password)
        print(" (définissez ADMIN_PASSWORD pour le choisir vous-même ; changez-le après connexion)")
        print("=" * 70)
    plant_id = None
    if os.environ.get("SEED_BASE_DATA", "1") != "0":
        cur = conn.execute("INSERT INTO plants(code,name,city) VALUES(?,?,?)", BASE_PLANT)
        plant_id = cur.lastrowid
        cur = conn.execute("INSERT INTO lines(plant_id,code,name) VALUES(?,?,?)", (plant_id,) + BASE_LINE)
        line_id = cur.lastrowid
        for pos, (ref, name) in enumerate(EQUIPMENT_COBOMI, 1):
            full = name if "CSD PET 4" in name else name + " - CSD PET 4"
            conn.execute("INSERT INTO equipment(line_id,sga_id,reference,name,position) VALUES(?,?,?,?,?)",
                         (line_id, None, ref, full, pos))
    conn.execute("INSERT INTO users(name,username,email,role,password_hash,plant_id,permissions,active) VALUES(?,?,?,?,?,?,?,1)",
                 (ADMIN_NAME, ADMIN_USER, ADMIN_EMAIL, "Administrateur", generate_password_hash(password), plant_id,
                  json.dumps(ROLE_DEFAULTS["Administrateur"])))
    conn.execute("UPDATE users SET must_change_password=1 WHERE username=?", (ADMIN_USER,))




# ---------------------------------------------------------------------------
# Authentification, droits et périmètre
# ---------------------------------------------------------------------------
@app.before_request
def load_user():
    g.user = None
    g.perms = set()
    if request.endpoint in ("static", "healthz"):
        return
    uid = session.get("uid")
    if uid:
        u = q("SELECT * FROM users WHERE id=? AND active=1", (uid,), one=True)
        if u:
            g.user = u
            g.perms = set(PERM_KEYS) if u["role"] == "Administrateur" else set(json.loads(u["permissions"] or "[]"))
    if g.user is None and request.endpoint not in ("login", "static", "service_worker", "healthz"):
        return redirect(url_for("login", next=request.path))
    if g.user is not None and g.user["must_change_password"] and request.endpoint not in (
            "account", "logout", "static", "service_worker", "healthz", "uploads"):
        return redirect(url_for("account"))


def can(perm):
    return g.user is not None and (g.user["role"] == "Administrateur" or perm in g.perms)


def is_admin():
    return g.user is not None and g.user["role"] == "Administrateur"


def require(perm):
    def deco(f):
        @wraps(f)
        def wrapper(*a, **kw):
            if not can(perm):
                abort(403)
            return f(*a, **kw)
        return wrapper
    return deco


def scope_sga_ids(u=None):
    """None = toutes les SGA ; sinon liste d'identifiants."""
    u = u or g.user
    if u["role"] in GLOBAL_ROLES:
        return None
    if u["role"] in PLANT_ROLES:
        return [r["id"] for r in q("SELECT s.id FROM sga s JOIN lines l ON l.id=s.line_id WHERE l.plant_id=?", (u["plant_id"],))]
    if u["role"] in CHEF_ROLES:
        # ses propres SGA + les SGA des membres de son équipe
        return [r["sga_id"] for r in q("""SELECT sga_id FROM user_sga WHERE user_id=? UNION
                                          SELECT us.sga_id FROM user_sga us JOIN users m ON m.id=us.user_id
                                          WHERE m.supervisor_id=? AND m.active=1""", (u["id"], u["id"]))]
    return [r["sga_id"] for r in q("SELECT sga_id FROM user_sga WHERE user_id=?", (u["id"],))]


def in_clause(col, ids):
    if ids is None:
        return "1=1", []
    if not ids:
        return "1=0", []
    return "%s IN (%s)" % (col, ",".join("?" * len(ids))), list(ids)


def scoped_sgas():
    clause, args = in_clause("s.id", scope_sga_ids())
    return q("SELECT s.*, l.name line_name, l.plant_id, p.name plant_name FROM sga s JOIN lines l ON l.id=s.line_id "
             "JOIN plants p ON p.id=l.plant_id WHERE " + clause + " ORDER BY s.name", args)


def scoped_equipment():
    clause, args = in_clause("e.sga_id", scope_sga_ids())
    return q("SELECT e.*, s.name sga_name FROM equipment e LEFT JOIN sga s ON s.id=e.sga_id WHERE " + clause +
             " ORDER BY e.position, e.name", args)


def team_members(u=None, sga_id=None):
    """Membres suivis : pour un chef d'équipe, les opérateurs / techniciens rattachés à lui ; pour un leader SGA, sa SGA."""
    u = u or g.user
    if u["role"] in CHEF_ROLES:
        sql = "SELECT DISTINCT m.* FROM users m WHERE m.active=1 AND m.supervisor_id=?"
        args = [u["id"]]
        if sga_id:
            sql += " AND m.id IN (SELECT user_id FROM user_sga WHERE sga_id=?)"
            args.append(int(sga_id))
        return q(sql + " ORDER BY m.role, m.name", args)
    ids = scope_sga_ids(u)
    if sga_id:
        ids = [int(sga_id)] if ids is None or int(sga_id) in ids else []
    clause, args = in_clause("us.sga_id", ids)
    roles = FIELD_ROLES
    if u["role"] == "Chef d'équipe Production":
        roles = ("Opérateur",)
    elif u["role"] == "Chef d'équipe Maintenance":
        roles = ("Technicien",)
    return q("SELECT DISTINCT us_u.* FROM users us_u JOIN user_sga us ON us.user_id=us_u.id WHERE us_u.active=1 AND " + clause +
             " AND us_u.role IN (%s) ORDER BY us_u.role, us_u.name" % ",".join("?" * len(roles)), args + list(roles))


def can_act_on_tag(tag):
    if can("tags_all"):
        return True
    if not can("tags_move"):
        return False
    ids = scope_sga_ids()
    return ids is None or tag["sga_id"] in ids


def user_sga_names(user_id):
    return ", ".join(r["name"] for r in q("SELECT s.name FROM sga s JOIN user_sga us ON us.sga_id=s.id WHERE us.user_id=? ORDER BY s.name", (user_id,)))


@app.context_processor
def inject():
    ctx = dict(can=can, is_admin=is_admin, me=g.get("user"), storage_warning=storage_warning(), TAG_TYPES=TAG_TYPES, TAG_TYPE=TAG_TYPE, STATUSES=STATUSES,
               STATUS_INFO=STATUS_INFO, ANOMALY_TYPES=ANOMALY_TYPES, FREQUENCIES=FREQUENCIES, MACHINE_STATES=MACHINE_STATES,
               DOC_TYPES=DOC_TYPES, ROLES=ROLES, PERMS=PERMS, PDR_COLORS=PDR_COLORS, today_str=today().strftime("%d/%m/%Y"), today_iso=today().isoformat(),
               CIL_TYPES=CIL_TYPES, CIL=CIL)
    if g.get("user"):
        ctx["me_sgas"] = user_sga_names(g.user["id"])
        ctx["nav_cm"] = q("SELECT COUNT(*) n FROM countermeasures WHERE status='À intégrer'", one=True)["n"] if can("checklists_manage") else 0
        ctx["nav_tasks"] = q("SELECT COUNT(*) n FROM tasks WHERE assigned_to=? AND status IN ('À réaliser','En cours') AND due_date<='" + today().isoformat() + "'",
                             (g.user["id"],), one=True)["n"]
    return ctx


@app.errorhandler(403)
def forbidden(e):
    return render_template("message.html", title="Accès refusé",
                           message="Votre profil n'a pas accès à cette page. Contactez l'administrateur pour obtenir le droit."), 403


@app.errorhandler(404)
def not_found(e):
    return render_template("message.html", title="Page introuvable", message="L'élément demandé n'existe pas."), 404


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        u = q("SELECT * FROM users WHERE username=? AND active=1", (request.form.get("username", "").strip(),), one=True)
        if u and check_password_hash(u["password_hash"], request.form.get("password", "")):
            session.clear()
            session.permanent = True
            session["uid"] = u["id"]
            if u["must_change_password"]:
                flash("Première connexion : choisissez votre mot de passe personnel pour continuer.", "error")
                return redirect(url_for("account"))
            nxt = request.args.get("next") or url_for("home")
            return redirect(nxt if nxt.startswith("/") and not nxt.startswith(("//", "/\\")) else url_for("home"))
        flash("Identifiant ou mot de passe incorrect.", "error")
    return render_template("login.html")


@app.route("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


# ---------------------------------------------------------------------------
# Indicateurs
# ---------------------------------------------------------------------------
def pct(a, b):
    return round(100 * a / b) if b else 0


def week_bounds(d=None):
    d = d or today()
    start = d - dt.timedelta(days=d.weekday())
    return start, start + dt.timedelta(days=6)


def period_since(key):
    T = today()
    return {"semaine": week_bounds()[0], "mois": T - dt.timedelta(days=30), "trimestre": T - dt.timedelta(days=90)}.get(key)


def user_kpis(user_id, since=None):
    since = since or week_bounds()[0]
    rows = q("SELECT status FROM tasks WHERE assigned_to=? AND due_date>=? AND due_date<=?", (user_id, since.isoformat(), today().isoformat()))
    planned = len(rows)
    done = sum(1 for r in rows if r["status"] in TASK_DONE)
    open_tags = q("SELECT COUNT(*) n FROM tags WHERE (treated_by=? OR (observed_by=? AND treated_by IS NULL)) AND status IN ('Nouveau','En cours')",
                  (user_id, user_id), one=True)["n"]
    emitted = q("SELECT COUNT(*) n FROM tags WHERE observed_by=? AND created_at>=?", (user_id, since.isoformat()), one=True)["n"]
    s5 = q("SELECT b.score_5s FROM sga_board b JOIN user_sga us ON us.sga_id=b.sga_id WHERE us.user_id=? LIMIT 1", (user_id,), one=True)
    return {"planned": planned, "done": done, "rate": pct(done, planned), "open_tags": open_tags, "emitted": emitted,
            "score_5s": s5["score_5s"] if s5 and s5["score_5s"] is not None else None}


def dashboard_data(f):
    """f : dict de filtres (plant, line, sga, chef, user, period)."""
    T = today()
    since = period_since(f.get("period") or "mois")
    clauses_t, args_t = [], []
    clauses_k, args_k = [], []
    sc, sa = in_clause("t.sga_id", scope_sga_ids())
    clauses_t.append(sc); args_t += sa
    sc, sa = in_clause("e.sga_id", scope_sga_ids())
    clauses_k.append(sc); args_k += sa
    for col, key in (("l.plant_id", "plant"), ("e.line_id", "line"), ("e.sga_id", "sga")):
        if f.get(key):
            clauses_t.append(col + "=?"); args_t.append(f[key])
            clauses_k.append(col + "=?"); args_k.append(f[key])
    users = None
    if f.get("user"):
        users = [int(f["user"])]
    elif f.get("chef"):
        chef = q("SELECT * FROM users WHERE id=?", (f["chef"],), one=True)
        users = [m["id"] for m in team_members(chef)] if chef else []
    if users is not None:
        c1, a1 = in_clause("t.observed_by", users)
        c2, a2 = in_clause("t.treated_by", users)
        clauses_t.append("(%s OR %s)" % (c1, c2)); args_t += a1 + a2
        c3, a3 = in_clause("k.assigned_to", users)
        clauses_k.append(c3); args_k += a3
    base_t = " FROM tags t JOIN equipment e ON e.id=t.equipment_id JOIN lines l ON l.id=e.line_id WHERE " + " AND ".join(clauses_t)
    base_k = (" FROM tasks k JOIN checklists c ON c.id=k.checklist_id JOIN equipment e ON e.id=c.equipment_id "
              "JOIN lines l ON l.id=e.line_id WHERE " + " AND ".join(clauses_k))
    per_t, per_ta = "", []
    per_k, per_ka = " AND k.due_date<=?", [T.isoformat()]
    if since:
        per_t, per_ta = " AND t.created_at>=?", [since.isoformat()]
        per_k += " AND k.due_date>=?"; per_ka.append(since.isoformat())
    tags = q("SELECT t.*, e.name eq_name" + base_t + per_t, args_t + per_ta)
    tasks = q("SELECT k.status, k.duration_min, c.duration_min std" + base_k + per_k, args_k + per_ka)
    total = len(tags)
    treated = sum(1 for t in tags if t["status"] in ("Traité", "Contre-mesure"))
    by_type = [{"t": tt, "n": sum(1 for t in tags if t["type"] == tt["key"]),
                "treated": sum(1 for t in tags if t["type"] == tt["key"] and t["status"] in ("Traité", "Contre-mesure"))} for tt in TAG_TYPES]
    by_status = [{"s": s, "n": sum(1 for t in tags if t["status"] == s)} for s in STATUSES]
    by_anomaly = sorted([{"a": a, "n": sum(1 for t in tags if t["anomaly_type"] == a)} for a in ANOMALY_TYPES], key=lambda x: -x["n"])
    planned = len(tasks)
    done = sum(1 for k in tasks if k["status"] in TASK_DONE)
    to_confirm = sum(1 for k in tasks if k["status"] == "Réalisée") + sum(
        1 for t in tags if t["status"] in ("Traité", "Contre-mesure") and not t["confirmed_at"])
    # Transfert RCM -> AM
    tr_clause, tr_args = in_clause("e.sga_id", scope_sga_ids())
    extra, extra_a = "", []
    for col, key in (("l.plant_id", "plant"), ("e.line_id", "line"), ("e.sga_id", "sga")):
        if f.get(key):
            extra += " AND %s=?" % col; extra_a.append(f[key])
    tr = q("SELECT SUM(c.transferred) tr, SUM(CASE WHEN c.activity_type='RCM' THEN 1 ELSE 0 END) rcm FROM checklists c "
           "JOIN equipment e ON e.id=c.equipment_id JOIN lines l ON l.id=e.line_id WHERE c.active=1 AND " + tr_clause + extra,
           tr_args + extra_a, one=True)
    transfer = pct(tr["tr"] or 0, (tr["tr"] or 0) + (tr["rcm"] or 0))
    # Audits
    au_clause, au_args = in_clause("a.sga_id", scope_sga_ids())
    au = q("SELECT AVG(a.score) s FROM audits a JOIN sga s ON s.id=a.sga_id JOIN lines l ON l.id=s.line_id WHERE " + au_clause +
           (" AND l.plant_id=?" if f.get("plant") else "") + (" AND s.line_id=?" if f.get("line") else "") +
           (" AND a.sga_id=?" if f.get("sga") else ""),
           au_args + [x for x in (f.get("plant"), f.get("line"), f.get("sga")) if x], one=True)
    # SGA ranking
    sga_rows = []
    for s in scoped_sgas():
        if f.get("sga") and str(s["id"]) != str(f["sga"]):
            continue
        if f.get("line") and str(s["line_id"]) != str(f["line"]):
            continue
        if f.get("plant") and str(s["plant_id"]) != str(f["plant"]):
            continue
        st = [t for t in tags if t["sga_id"] == s["id"]]
        tk = q("SELECT k.status FROM tasks k JOIN checklists c ON c.id=k.checklist_id JOIN equipment e ON e.id=c.equipment_id "
               "WHERE e.sga_id=? AND k.due_date<=?" + (" AND k.due_date>=?" if since else ""),
               [s["id"], T.isoformat()] + ([since.isoformat()] if since else []))
        last_audit = q("SELECT score, level FROM audits WHERE sga_id=? ORDER BY audit_date DESC, id DESC LIMIT 1", (s["id"],), one=True)
        sga_rows.append({"s": s, "open": sum(1 for t in st if t["status"] in ("Nouveau", "En cours")), "total": len(st),
                         "rate": pct(sum(1 for k in tk if k["status"] in TASK_DONE), len(tk)),
                         "audit": last_audit})
    # Tendance 8 semaines
    trend = []
    ws, _ = week_bounds()
    for i in range(7, -1, -1):
        a = ws - dt.timedelta(weeks=i)
        b = a + dt.timedelta(days=7)
        created = q("SELECT COUNT(*) n" + base_t + " AND t.created_at>=? AND t.created_at<?", args_t + [a.isoformat(), b.isoformat()], one=True)["n"]
        resolved = q("SELECT COUNT(*) n" + base_t + " AND t.resolved_at>=? AND t.resolved_at<?", args_t + [a.isoformat(), b.isoformat()], one=True)["n"]
        trend.append({"label": a.strftime("%d/%m"), "created": created, "resolved": resolved})
    return {"total": total, "treated": treated, "treat_rate": pct(treated, total), "by_type": by_type, "by_status": by_status,
            "by_anomaly": by_anomaly, "planned": planned, "done": done, "real_rate": pct(done, planned),
            "to_confirm": to_confirm, "transfer": transfer, "audit_avg": round(au["s"]) if au and au["s"] else None,
            "sga_rows": sga_rows, "trend": trend, "trend_max": max([1] + [max(x["created"], x["resolved"]) for x in trend]),
            "open": sum(1 for t in tags if t["status"] in ("Nouveau", "En cours"))}


# ---------------------------------------------------------------------------
# Accueil
# ---------------------------------------------------------------------------
@app.route("/")
def home():
    u = g.user
    if u["role"] in FIELD_ROLES:
        return home_field()
    if u["role"] in TEAM_ROLES:
        return home_team()
    return dashboard(home=True)


def task_rows(where, args):
    return q("""SELECT k.*, c.name cl_name, COALESCE((SELECT SUM(ti.duration_min) FROM task_items ti WHERE ti.task_id=k.id), c.duration_min) std,
               COALESCE((SELECT GROUP_CONCAT(DISTINCT ti.organe) FROM task_items ti WHERE ti.task_id=k.id AND ti.organe!=''), c.organe) organe,
               COALESCE((SELECT GROUP_CONCAT(DISTINCT ti.frequency) FROM task_items ti WHERE ti.task_id=k.id), c.frequency) frequency, e.name eq_name, u.name user_name,
               (SELECT COUNT(*) FROM task_items i WHERE i.task_id=k.id) n_items,
               (SELECT COUNT(*) FROM task_items i WHERE i.task_id=k.id AND i.ok IS NOT NULL) n_done
               FROM tasks k JOIN checklists c ON c.id=k.checklist_id JOIN equipment e ON e.id=c.equipment_id
               LEFT JOIN users u ON u.id=k.assigned_to WHERE """ + where, args)


def home_field():
    u = g.user
    k = user_kpis(u["id"])
    tasks = task_rows("k.assigned_to=? AND k.due_date<=? AND (k.status IN ('À réaliser','En cours') OR k.due_date>=?) ORDER BY "
                      "CASE k.status WHEN 'En cours' THEN 0 WHEN 'À réaliser' THEN 1 ELSE 2 END, k.due_date LIMIT 12",
                      (u["id"], today().isoformat(), week_bounds()[0].isoformat()))
    clause, args = in_clause("t.sga_id", scope_sga_ids())
    tags = q("SELECT t.*, e.name eq_name FROM tags t JOIN equipment e ON e.id=t.equipment_id WHERE " + clause +
             " AND (t.observed_by=? OR t.treated_by=? OR t.status IN ('Nouveau','En cours')) ORDER BY t.created_at DESC LIMIT 6",
             args + [u["id"], u["id"]])
    docs = q("SELECT d.*, e.name eq_name FROM documents d JOIN equipment e ON e.id=d.equipment_id WHERE " +
             in_clause("e.sga_id", scope_sga_ids())[0] + " ORDER BY d.doc_type, d.title LIMIT 6", in_clause("e.sga_id", scope_sga_ids())[1])
    return render_template("home_field.html", k=k, tasks=tasks, tags=tags, docs=docs)


def home_team():
    u = g.user
    sga_id = request.args.get("sga", "")
    period = request.args.get("period", "semaine")
    member_id = request.args.get("member", "")
    since = period_since(period)
    members = team_members(u, sga_id or None)
    member_ids = [m["id"] for m in members]
    shown = [m for m in members if not member_id or str(m["id"]) == member_id]
    perf = []
    for m in shown:
        k = user_kpis(m["id"], since)
        perf.append({"m": m, "k": k, "sga": user_sga_names(m["id"])})
    ids = [int(sga_id)] if sga_id else scope_sga_ids()
    clause, args = in_clause("t.sga_id", ids)
    tags = q("SELECT t.status, t.confirmed_at FROM tags t WHERE " + clause + " AND t.created_at>=?", args + [since.isoformat()])
    to_confirm_tasks, to_confirm_tags = [], []
    if can("confirm") and member_ids:
        mc, ma = in_clause("k.assigned_to", member_ids)
        to_confirm_tasks = task_rows(mc + " AND k.status='Réalisée' ORDER BY k.completed_at DESC", ma)
        tc, ta = in_clause("t.treated_by", member_ids)
        to_confirm_tags = q("SELECT t.*, e.name eq_name, u.name user_name FROM tags t JOIN equipment e ON e.id=t.equipment_id "
                            "LEFT JOIN users u ON u.id=t.treated_by WHERE " + tc +
                            " AND t.status IN ('Traité','Contre-mesure') AND t.confirmed_at IS NULL ORDER BY t.resolved_at DESC", ta)
    planned = sum(p["k"]["planned"] for p in perf)
    done = sum(p["k"]["done"] for p in perf)
    kpi = {"tags": len(tags), "open": sum(1 for t in tags if t["status"] == "En cours"),
           "new": sum(1 for t in tags if t["status"] == "Nouveau"),
           "to_confirm": len(to_confirm_tasks) + len(to_confirm_tags),
           "treated": sum(1 for t in tags if t["status"] in ("Traité", "Contre-mesure")), "rate": pct(done, planned)}
    my_tasks = task_rows("k.assigned_to=? AND k.status IN ('À réaliser','En cours') ORDER BY k.due_date", (u["id"],))
    to_validate = SAP["pending_validation"](u) if can("tag_validate") else []
    return render_template("home_team.html", to_validate=to_validate, sap_linked=SAP["has_token"](u["id"]), perf=perf, kpi=kpi, members=members, sgas=scoped_sgas(), sel_sga=sga_id,
                           sel_member=member_id, period=period, to_confirm_tasks=to_confirm_tasks,
                           to_confirm_tags=to_confirm_tags, my_tasks=my_tasks)


@app.route("/dashboard")
def dashboard(home=False):
    if not home and not (can("dashboard_global") or g.user["role"] in TEAM_ROLES):
        abort(403)
    f = {k: request.args.get(k, "") for k in ("plant", "line", "sga", "chef", "user", "period")}
    f["period"] = f["period"] or "mois"
    data = dashboard_data(f)
    sc, sa = in_clause("us.sga_id", scope_sga_ids())
    users = q("SELECT DISTINCT u.* FROM users u JOIN user_sga us ON us.user_id=u.id WHERE u.active=1 AND " + sc +
              " AND u.role IN ('Opérateur','Technicien') ORDER BY u.name", sa)
    chefs = q("SELECT * FROM users WHERE active=1 AND role IN (?,?,?) ORDER BY name", TEAM_ROLES)
    plants = q("SELECT * FROM plants ORDER BY name")
    lines = q("SELECT * FROM lines ORDER BY name")
    to_validate = SAP["pending_validation"](g.user) if (home and can("tag_validate")) else []
    return render_template("dashboard.html", to_validate=to_validate, d=data, f=f, sgas=scoped_sgas(), users=users, chefs=chefs, plants=plants,
                           lines=lines, is_home=home)


# ---------------------------------------------------------------------------
# Tâches (exécution des checklists)
# ---------------------------------------------------------------------------
@app.route("/tasks")
def my_tasks():
    u = g.user
    open_tasks = task_rows("k.assigned_to=? AND k.status IN ('À réaliser','En cours') AND k.due_date<=? ORDER BY k.due_date",
                           (u["id"], today().isoformat()))
    upcoming = task_rows("k.assigned_to=? AND k.status='À réaliser' AND k.due_date>? ORDER BY k.due_date", (u["id"], today().isoformat()))
    history = task_rows("k.assigned_to=? AND k.status NOT IN ('À réaliser','En cours') ORDER BY k.due_date DESC, k.id DESC LIMIT 40", (u["id"],))
    return render_template("tasks.html", open_tasks=open_tasks, upcoming=upcoming, history=history, k=user_kpis(u["id"]))


def load_task(task_id):
    t = q("""SELECT k.*, c.name cl_name, COALESCE((SELECT SUM(ti.duration_min) FROM task_items ti WHERE ti.task_id=k.id), c.duration_min) std,
               COALESCE((SELECT GROUP_CONCAT(DISTINCT ti.organe) FROM task_items ti WHERE ti.task_id=k.id AND ti.organe!=''), c.organe) organe,
               COALESCE((SELECT GROUP_CONCAT(DISTINCT ti.frequency) FROM task_items ti WHERE ti.task_id=k.id), c.frequency) frequency, c.sop_url, c.machine_state, c.people,
             c.equipment_id, c.code cl_code, c.countermeasure, e.name eq_name, e.sga_id, u.name user_name
             FROM tasks k JOIN checklists c ON c.id=k.checklist_id JOIN equipment e ON e.id=c.equipment_id
             LEFT JOIN users u ON u.id=k.assigned_to WHERE k.id=?""", (task_id,), one=True)
    if not t:
        abort(404)
    return t


def can_view_task(t):
    if t["assigned_to"] == g.user["id"] or is_admin() or can("checklists_manage"):
        return True
    ids = scope_sga_ids()
    return ids is None or t["sga_id"] in ids


@app.route("/tasks/<int:task_id>")
def task_run(task_id):
    t = load_task(task_id)
    if not can_view_task(t):
        abort(403)
    items = q("SELECT i.*, tg.tag_code FROM task_items i LEFT JOIN tags tg ON tg.id=i.tag_id WHERE i.task_id=? ORDER BY i.position", (task_id,))
    editable = t["assigned_to"] == g.user["id"] and t["status"] in TASK_OPEN and can("tasks_exec")
    docs = q("SELECT * FROM documents WHERE equipment_id=? ORDER BY doc_type", (t["equipment_id"],))
    return render_template("task_run.html", t=t, items=items, editable=editable, docs=docs,
                           all_done=all(i["ok"] is not None for i in items))


@app.route("/tasks/<int:task_id>/item/<int:item_id>", methods=["POST"])
def task_item(task_id, item_id):
    t = load_task(task_id)
    if t["assigned_to"] != g.user["id"] or t["status"] not in TASK_OPEN:
        abort(403)
    ex("UPDATE task_items SET ok=1, validated_at=? WHERE id=? AND task_id=?", (now(), item_id, task_id))
    if t["status"] == "À réaliser":
        ex("UPDATE tasks SET status='En cours', started_at=? WHERE id=?", (now(), task_id))
    return redirect(url_for("task_run", task_id=task_id))


@app.route("/tasks/<int:task_id>/item/<int:item_id>/reset", methods=["POST"])
def task_item_reset(task_id, item_id):
    t = load_task(task_id)
    if t["assigned_to"] != g.user["id"] or t["status"] not in TASK_OPEN:
        abort(403)
    ex("UPDATE task_items SET ok=NULL, validated_at=NULL WHERE id=? AND task_id=? AND tag_id IS NULL", (item_id, task_id))
    return redirect(url_for("task_run", task_id=task_id))


@app.route("/tasks/<int:task_id>/finish", methods=["POST"])
def task_finish(task_id):
    t = load_task(task_id)
    if t["assigned_to"] != g.user["id"] or t["status"] not in TASK_OPEN:
        abort(403)
    pending = q("SELECT COUNT(*) n FROM task_items WHERE task_id=? AND ok IS NULL", (task_id,), one=True)["n"]
    if pending:
        flash("Tous les points doivent être validés avant de terminer la tâche.", "error")
        return redirect(url_for("task_run", task_id=task_id))
    started = t["started_at"] or now()
    try:
        dur = int((dt.datetime.now() - dt.datetime.strptime(started, "%Y-%m-%d %H:%M")).total_seconds() // 60)
    except ValueError:
        dur = None
    ex("UPDATE tasks SET status='Réalisée', completed_at=?, started_at=?, duration_min=?, reject_comment=NULL WHERE id=?",
       (now(), started, max(dur or 1, 1), task_id))
    flash("Tâche terminée : elle est envoyée au chef d'équipe pour confirmation.", "success")
    return redirect(url_for("home"))


@app.route("/confirm/task/<int:task_id>", methods=["POST"])
@require("confirm")
def confirm_task(task_id):
    t = load_task(task_id)
    if request.form.get("action") == "reject":
        comment = request.form.get("comment", "").strip()
        if not comment:
            flash("Un commentaire est obligatoire pour rejeter.", "error")
            return redirect(request.referrer or url_for("home"))
        ex("UPDATE tasks SET status='À réaliser', reject_comment=?, completed_at=NULL WHERE id=?", (comment, task_id))
        ex("UPDATE task_items SET ok=NULL, validated_at=NULL WHERE task_id=? AND tag_id IS NULL", (task_id,))
        flash("Tâche rejetée : elle revient chez %s." % t["user_name"], "success")
    else:
        ex("UPDATE tasks SET status='Confirmée', confirmed_by=?, confirmed_at=? WHERE id=?", (g.user["id"], now(), task_id))
        flash("Tâche confirmée.", "success")
    return redirect(request.referrer or url_for("home"))


@app.route("/confirm/tag/<int:tag_id>", methods=["POST"])
@require("confirm")
def confirm_tag(tag_id):
    tag = load_tag(tag_id)
    if request.form.get("action") == "reject":
        comment = request.form.get("comment", "").strip() or "Traitement non conforme"
        ex("UPDATE tags SET status='En cours', resolved_at=NULL, updated_at=? WHERE id=?", (now(), tag_id))
        ex("INSERT INTO tag_history(tag_id,user_id,from_status,to_status,comment,at) VALUES(?,?,?,?,?,?)",
           (tag_id, g.user["id"], tag["status"], "En cours", "Rejeté : " + comment, now()))
        flash("TAG renvoyé en cours de traitement.", "success")
    else:
        ex("UPDATE tags SET confirmed_by=?, confirmed_at=? WHERE id=?", (g.user["id"], now(), tag_id))
        ex("INSERT INTO tag_history(tag_id,user_id,from_status,to_status,comment,at) VALUES(?,?,?,?,?,?)",
           (tag_id, g.user["id"], tag["status"], tag["status"], "Traitement confirmé", now()))
        flash("Traitement du TAG confirmé.", "success")
    return redirect(request.referrer or url_for("home"))


# ---------------------------------------------------------------------------
# TAGs
# ---------------------------------------------------------------------------
def load_tag(tag_id):
    t = q("""SELECT t.*, e.name eq_name, e.reference eq_ref, s.name sga_name, l.name line_name, p.name plant_name,
             ob.name observed_name, tr.name treated_name, cf.name confirmed_name
             FROM tags t JOIN equipment e ON e.id=t.equipment_id LEFT JOIN sga s ON s.id=t.sga_id
             JOIN lines l ON l.id=e.line_id JOIN plants p ON p.id=l.plant_id
             LEFT JOIN users ob ON ob.id=t.observed_by LEFT JOIN users tr ON tr.id=t.treated_by
             LEFT JOIN users cf ON cf.id=t.confirmed_by WHERE t.id=?""", (tag_id,), one=True)
    if not t:
        abort(404)
    return t


def tag_query(extra="", args=()):
    clause, sargs = in_clause("t.sga_id", scope_sga_ids())
    return q("""SELECT t.*, e.name eq_name, s.name sga_name, ob.name observed_name, tr.name treated_name
               FROM tags t JOIN equipment e ON e.id=t.equipment_id LEFT JOIN sga s ON s.id=t.sga_id
               JOIN lines l ON l.id=e.line_id
               LEFT JOIN users ob ON ob.id=t.observed_by LEFT JOIN users tr ON tr.id=t.treated_by
               WHERE """ + clause + extra, sargs + list(args))


def tag_filters():
    f = {k: request.args.get(k, "") for k in ("plant", "line", "sga", "equipment", "type", "status", "anomaly", "search")}
    extra, args = "", []
    for col, key in (("l.plant_id", "plant"), ("e.line_id", "line"), ("t.sga_id", "sga"), ("t.equipment_id", "equipment"),
                     ("t.type", "type"), ("t.status", "status"), ("t.anomaly_type", "anomaly")):
        if f[key]:
            extra += " AND %s=?" % col
            args.append(f[key])
    if f["search"]:
        extra += " AND (t.tag_code LIKE ? OR t.description LIKE ? OR t.organe LIKE ? OR e.name LIKE ?)"
        args += ["%" + f["search"] + "%"] * 4
    return f, extra, args


def assignable_users(sga_id=None):
    clause, args = in_clause("us.sga_id", [sga_id] if sga_id else scope_sga_ids())
    rows = q("SELECT DISTINCT u.* FROM users u JOIN user_sga us ON us.user_id=u.id WHERE u.active=1 AND " + clause +
             " ORDER BY CASE u.role WHEN 'Opérateur' THEN 0 WHEN 'Technicien' THEN 1 ELSE 2 END, u.name", args)
    if not rows:
        rows = q("SELECT * FROM users WHERE active=1 AND role IN ('Opérateur','Technicien') ORDER BY name")
    return rows


@app.route("/tags")
def tags_board():
    f, extra, args = tag_filters()
    rows = tag_query(extra + " ORDER BY t.created_at DESC", args)
    board = {tt["key"]: {s: [] for s in STATUSES} for tt in TAG_TYPES}
    for t in rows:
        if t["type"] in board and t["status"] in STATUSES:
            board[t["type"]][t["status"]].append(t)
    sel_sga = q("SELECT s.*, l.name line_name, p.name plant_name FROM sga s JOIN lines l ON l.id=s.line_id JOIN plants p ON p.id=l.plant_id WHERE s.id=?",
                (f["sga"],), one=True) if f["sga"] else None
    return render_template("tags_board.html", board=board, f=f, sgas=scoped_sgas(), equipment=scoped_equipment(),
                           sel_sga=sel_sga, plants=q("SELECT * FROM plants"), lines=q("SELECT * FROM lines"),
                           users=active_users(), count=len(rows))


@app.route("/tags/registre")
def tags_list():
    f, extra, args = tag_filters()
    rows = tag_query(extra + " ORDER BY t.created_at DESC", args)
    return render_template("tags_list.html", rows=rows, f=f, sgas=scoped_sgas(), equipment=scoped_equipment(),
                           plants=q("SELECT * FROM plants"), lines=q("SELECT * FROM lines"), users=assignable_users())


@app.route("/tags/registre/export.xlsx")
def tags_export():
    """Extraction Excel du registre des TAGs (administrateur), avec les filtres de la page."""
    if not is_admin():
        abort(403)
    import io
    from flask import send_file
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    f, extra, args = tag_filters()
    rows = tag_query(extra + " ORDER BY t.created_at DESC", args)
    names = {u["id"]: u["name"] for u in q("SELECT id, name FROM users")}

    def when(v):
        v = (v or "")[:19]
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%dT%H:%M"):
            try:
                return dt.datetime.strptime(v, fmt)
            except ValueError:
                pass
        try:
            return dt.datetime.strptime(v, "%Y-%m-%d").date()
        except ValueError:
            return v or None

    cols = [("Code", 12), ("Type", 9), ("Type d'anomalie", 24), ("SGA", 18), ("Machine", 32), ("Organe", 20),
            ("Description", 50), ("Découverte par", 20), ("Date d'observation", 18), ("Date de création", 18),
            ("Statut", 14), ("Validé par", 20), ("Date de validation", 18), ("N° avis SAP", 14), ("N° OT", 14),
            ("Traitée par", 20), ("Début traitement", 18), ("Date de résolution", 18), ("Contre-mesure", 40), ("Photo", 8)]
    wb = Workbook()
    ws = wb.active
    ws.title = "Registre des TAGs"
    ws.append([c for c, _ in cols])
    for t in rows:
        ws.append([t["tag_code"], TAG_TYPE.get(t["type"], {}).get("label", t["type"]), t["anomaly_type"], t["sga_name"],
                   t["eq_name"], t["organe"], t["description"], t["observed_name"], when(t["observed_at"]), when(t["created_at"]),
                   STATUS_INFO.get(t["status"], {}).get("board", t["status"]), names.get(t["validated_by"]), when(t["validated_at"]),
                   t["sap_notif"], t["ot_number"] or t["sap_order"], t["treated_name"], when(t["treatment_start"]),
                   when(t["resolved_at"]), t["countermeasure"], "Oui" if t["photo"] else "Non"])
    head = PatternFill("solid", fgColor="C8102E")
    for i, (_, width) in enumerate(cols, 1):
        c = ws.cell(row=1, column=i)
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=c.column in (7, 19))
            if isinstance(c.value, dt.datetime):
                c.number_format = "dd/mm/yyyy hh:mm"
            elif isinstance(c.value, dt.date):
                c.number_format = "dd/mm/yyyy"
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    name = "registre_tags_%s.xlsx" % dt.datetime.now().strftime("%Y%m%d_%H%M")
    return send_file(out, as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def active_users():
    return q("SELECT id, name, role FROM users WHERE active=1 ORDER BY name")


def save_photo(file):
    if not (file and file.filename):
        return ""
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in (".jpg", ".jpeg", ".png", ".gif", ".webp", ".heic"):
        return ""
    name = secure_filename("tag_%s_%s%s" % (dt.datetime.now().strftime("%Y%m%d%H%M%S%f"), g.user["id"], ext))
    file.save(os.path.join(UPLOADS, name))
    pg_store.save_file(name, os.path.join(UPLOADS, name))
    return name


def needs_sap(t):
    return t["type"] in sap.SAP_TYPES


def clean_ref(v):
    return "".join(ch for ch in (v or "") if ch.isalnum() or ch in "-/")


@app.route("/tags/new", methods=["GET", "POST"])
@require("tags_create")
def tag_new():
    equipment = scoped_equipment() if not is_admin() else q(
        "SELECT e.*, s.name sga_name FROM equipment e LEFT JOIN sga s ON s.id=e.sga_id ORDER BY e.position, e.name")
    users = active_users()
    if request.method == "POST":
        form = request.form
        e = q("SELECT * FROM equipment WHERE id=?", (form.get("equipment_id") or 0,), one=True)
        errors = []
        if form.get("type") not in TAG_TYPE:
            errors.append("Choisissez le type d'étiquette (AM, RCM ou Safety).")
        if form.get("anomaly_type") not in ANOMALY_TYPES:
            errors.append("Choisissez le type d'anomalie.")
        if not e or e["id"] not in [x["id"] for x in equipment]:
            errors.append("Choisissez la machine.")
        if not form.get("description", "").strip():
            errors.append("Décrivez l'anomalie.")
        observer = q("SELECT id FROM users WHERE id=? AND active=1", (form.get("observed_by") or 0,), one=True)
        if not observer:
            errors.append("Choisissez la personne qui a découvert l'anomalie.")
        if errors:
            for er in errors:
                flash(er, "error")
            return render_template("tag_form.html", equipment=equipment, v=form, users=users)
        photo = save_photo(request.files.get("photo"))
        n = q("SELECT COALESCE(MAX(id),0)+1 n FROM tags", one=True)["n"]
        code = "TAG-%04d" % n
        while q("SELECT 1 FROM tags WHERE tag_code=?", (code,), one=True):
            n += 1
            code = "TAG-%04d" % n
        tid = ex("""INSERT INTO tags(tag_code,type,anomaly_type,equipment_id,sga_id,line_id,organe,description,photo,
                 observed_by,observed_at,created_at,ot_number,status,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,'', 'Nouveau', ?)""",
                 (code, form["type"], form["anomaly_type"], e["id"], e["sga_id"], e["line_id"], form.get("organe", "").strip(),
                  form["description"].strip(), photo, observer["id"], form.get("observed_at") or today().isoformat(), now(), now()))
        ex("INSERT INTO tag_history(tag_id,user_id,from_status,to_status,comment,at) VALUES(?,?,?,?,?,?)",
           (tid, g.user["id"], None, "Nouveau", "TAG créé", now()))
        item_id = form.get("task_item_id")
        if item_id:
            it = q("SELECT i.*, k.assigned_to, k.id task_id FROM task_items i JOIN tasks k ON k.id=i.task_id WHERE i.id=?", (item_id,), one=True)
            if it and it["assigned_to"] == g.user["id"]:
                ex("UPDATE task_items SET ok=0, validated_at=?, tag_id=? WHERE id=?", (now(), tid, item_id))
                ex("UPDATE tasks SET status='En cours', started_at=COALESCE(started_at, ?) WHERE id=? AND status='À réaliser'", (now(), it["task_id"]))
                flash("%s créé et rattaché au point non conforme." % code, "success")
                return redirect(url_for("task_run", task_id=it["task_id"]))
        msg = "%s créé : il apparaît dans « Anomalies découvertes »." % code
        if form["type"] in sap.SAP_TYPES:
            msg += " Il doit maintenant être validé par le chef d'équipe production (création de l'avis SAP)."
        flash(msg, "success")
        return redirect(url_for("tag_detail", tag_id=tid))
    v = {k: request.args.get(k, "") for k in ("type", "equipment_id", "organe", "description", "task_item_id")}
    v["observed_by"] = ""
    return render_template("tag_form.html", equipment=equipment, v=v, users=users)


@app.route("/tags/<int:tag_id>")
def tag_detail(tag_id):
    t = load_tag(tag_id)
    ids = scope_sga_ids()
    if not (is_admin() or can("tags_all") or ids is None or t["sga_id"] in ids or t["observed_by"] == g.user["id"]):
        abort(403)
    history = q("SELECT h.*, u.name user_name FROM tag_history h LEFT JOIN users u ON u.id=h.user_id WHERE h.tag_id=? ORDER BY h.at, h.id", (tag_id,))
    cms = q("""SELECT m.*, c.code cl_code, c.name cl_name FROM countermeasures m LEFT JOIN checklists c ON c.id=m.checklist_id
               WHERE m.tag_id=? ORDER BY m.id""", (tag_id,))
    pdr = q("SELECT * FROM pdr_requests WHERE tag_id=? ORDER BY id DESC", (tag_id,))
    idx = STATUSES.index(t["status"]) if t["status"] in STATUSES else 0
    targets = []
    blocked = needs_sap(t) and not t["sap_notif"]
    if can_act_on_tag(t) and not blocked:
        if is_admin():
            targets = [s for s in STATUSES if s != t["status"]]
        elif idx + 1 < len(STATUSES):
            targets = [STATUSES[idx + 1]]
        if "Contre-mesure" in targets and not can("countermeasure"):
            targets.remove("Contre-mesure")
    sap_payload = SAP["notif_payload"](t, g.user) if needs_sap(t) else None
    can_validate = can("tag_validate") and needs_sap(t) and not t["sap_notif"]
    sap_log = q("SELECT l.*, u.name user_name FROM sap_log l LEFT JOIN users u ON u.id=l.user_id WHERE l.tag_id=? ORDER BY l.id DESC LIMIT 10", (tag_id,))
    validated_name = q("SELECT name FROM users WHERE id=?", (t["validated_by"],), one=True) if t["validated_by"] else None
    return render_template("tag_detail.html", sap_payload=sap_payload, can_validate=can_validate, blocked=blocked,
                           sap_log=sap_log, validated_name=validated_name["name"] if validated_name else "",
                           t=t, history=history, cms=cms, pdr=pdr, targets=targets, users=active_users())


@app.route("/tags/<int:tag_id>/validate", methods=["POST"])
@require("tag_validate")
def tag_validate(tag_id):
    """Validation d'un TAG RCM / Safety par le chef d'équipe production : le N° de l'avis SAP créé est obligatoire."""
    t = load_tag(tag_id)
    ids = scope_sga_ids()
    team = [m["id"] for m in team_members()] if g.user["role"] in CHEF_ROLES else []
    other_chef = q("""SELECT s.id FROM users o JOIN users s ON s.id=o.supervisor_id
                      WHERE o.id=? AND s.role=? AND s.id!=?""", (t["observed_by"], g.user["role"], g.user["id"]), one=True)
    in_scope = (ids is None or t["sga_id"] in ids) and not (g.user["role"] in CHEF_ROLES and other_chef)
    if not (is_admin() or in_scope or t["observed_by"] in team):
        abort(403)
    back = url_for("tag_detail", tag_id=tag_id, validate=1)
    if not needs_sap(t):
        flash("Seuls les TAGs RCM et Safety passent par un avis SAP.", "error")
        return redirect(url_for("tag_detail", tag_id=tag_id))
    if t["sap_notif"]:
        flash("Ce TAG a déjà l'avis SAP %s." % t["sap_notif"], "error")
        return redirect(url_for("tag_detail", tag_id=tag_id))
    notif = clean_ref(request.form.get("sap_notif"))
    if not notif:
        flash("Le N° de l'avis SAP est obligatoire pour valider le TAG.", "error")
        return redirect(back)
    dup = q("SELECT tag_code FROM tags WHERE sap_notif=? AND id!=?", (notif, tag_id), one=True)
    if dup:
        flash("L'avis %s est déjà rattaché au %s. Vérifiez le numéro." % (notif, dup["tag_code"]), "error")
        return redirect(back)
    who = g.user["sap_user"] or g.user["name"]
    ex("""UPDATE tags SET validated_by=?, validated_at=?, sap_notif=?, sap_created_by=?, sap_sent_at=?, sap_status=?,
          updated_at=? WHERE id=?""", (g.user["id"], now(), notif, who, now(), "Avis créé", now(), tag_id))
    ex("INSERT INTO tag_history(tag_id,user_id,from_status,to_status,comment,at) VALUES(?,?,?,?,?,?)",
       (tag_id, g.user["id"], t["status"], t["status"], "Validé par %s — avis SAP %s" % (g.user["name"], notif), now()))
    sap.log(tag_id, "Validation / avis SAP", True, "Avis %s saisi par %s" % (notif, g.user["name"]))
    flash("%s validé : avis SAP %s. Le TAG peut maintenant passer en traitement." % (t["tag_code"], notif), "success")
    return redirect(request.form.get("next") or url_for("tag_detail", tag_id=tag_id))


@app.route("/tags/<int:tag_id>/move", methods=["POST"])
def tag_move(tag_id):
    t = load_tag(tag_id)
    wants_json = request.headers.get("X-Requested-With") == "fetch"
    target = request.form.get("to") or (request.get_json(silent=True) or {}).get("to")

    def fail(msg):
        if wants_json:
            return jsonify(ok=False, message=msg), 400
        flash(msg, "error")
        return redirect(request.referrer or url_for("tag_detail", tag_id=tag_id))

    if target not in STATUSES or target == t["status"]:
        return fail("Statut cible invalide.")
    if not can_act_on_tag(t):
        return fail("Vous n'avez pas le droit de déplacer ce TAG.")
    if needs_sap(t) and not t["sap_notif"] and target != "Nouveau":
        return fail("%s doit d'abord être validé par le chef d'équipe production, avec le N° de l'avis SAP." % t["tag_code"])
    cur_i, new_i = STATUSES.index(t["status"]), STATUSES.index(target)
    if not is_admin() and new_i != cur_i + 1:
        return fail("Un TAG descend d'un niveau à la fois. Seul l'administrateur peut le remonter.")
    if target == "Contre-mesure":
        if not can("countermeasure"):
            return fail("Le passage en contre-mesure est réservé à l'administrateur.")
        url = url_for("tag_countermeasure", tag_id=tag_id)
        return jsonify(ok=True, redirect=url) if wants_json else redirect(url)
    sets, vals = ["status=?", "updated_at=?", "confirmed_by=NULL", "confirmed_at=NULL"], [target, now()]
    comment = request.form.get("comment", "")
    ot = clean_ref(request.form.get("ot_number")) if needs_sap(t) else ""
    treated_form = request.form.get("treated_by")
    if treated_form and not q("SELECT 1 FROM users WHERE id=? AND active=1", (treated_form,), one=True):
        return fail("Personne inconnue.")
    if target == "En cours":
        sets += ["treated_by=?", "treatment_start=COALESCE(treatment_start, ?)"]
        vals += [treated_form or t["treated_by"] or g.user["id"], now()]
    if target == "Traité":
        treated_by = treated_form or t["treated_by"]
        if not treated_by:
            return fail("Indiquez la personne qui a traité l'anomalie.")
        if needs_sap(t) and not (ot or t["ot_number"]):
            return fail("Le N° d'OT est obligatoire pour passer un TAG %s en « Traité »." % TAG_TYPE[t["type"]]["label"])
        sets += ["resolved_at=?", "treated_by=?", "treatment_start=COALESCE(treatment_start, ?)"]
        vals += [now(), treated_by, now()]
    if ot:
        sets += ["ot_number=?", "sap_order=?"]
        vals += [ot, ot]
    if new_i < 2:
        sets.append("resolved_at=NULL")
    if new_i < 1:
        sets += ["treatment_start=NULL", "treated_by=NULL"]
    ex("UPDATE tags SET " + ", ".join(sets) + " WHERE id=?", vals + [tag_id])
    ex("INSERT INTO tag_history(tag_id,user_id,from_status,to_status,comment,at) VALUES(?,?,?,?,?,?)",
       (tag_id, g.user["id"], t["status"], target, comment or ("OT %s" % ot if ot else ""), now()))
    msg = "%s → %s" % (t["tag_code"], STATUS_INFO[target]["board"])
    if wants_json:
        return jsonify(ok=True, message=msg)
    flash(msg, "success")
    return redirect(request.referrer or url_for("tags_board"))


@app.route("/tags/<int:tag_id>/countermeasure", methods=["GET", "POST"])
@require("countermeasure")
def tag_countermeasure(tag_id):
    """La contre-mesure est enregistrée dans la liste des contre-mesures (elle n'est plus ajoutée automatiquement aux checklists)."""
    t = load_tag(tag_id)
    if needs_sap(t) and not t["sap_notif"]:
        flash("%s doit d'abord être validé par le chef d'équipe production (avis SAP)." % t["tag_code"], "error")
        return redirect(url_for("tag_detail", tag_id=tag_id))
    if request.method == "POST":
        f = request.form
        tasks = "\n".join(x.strip() for x in f.get("tasks", "").splitlines() if x.strip())
        errors = []
        if not f.get("countermeasure", "").strip():
            errors.append("Décrivez la contre-mesure.")
        if not tasks:
            errors.append("Indiquez au moins une tâche à réaliser.")
        if f.get("frequency") not in FREQUENCIES:
            errors.append("Choisissez la fréquence.")
        if errors:
            for er in errors:
                flash(er, "error")
            return render_template("countermeasure.html", t=t, v=f)
        ex("""INSERT INTO countermeasures(tag_id,equipment_id,organe,countermeasure,tasks,cil_type,tooling,duration_min,frequency,
              machine_state,people,sop_url,activity_type,status,created_by,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?, 'À intégrer', ?, ?)""",
           (tag_id, t["equipment_id"], f.get("organe", "").strip(), f["countermeasure"].strip(), tasks, f.get("cil_type") or None,
            f.get("tooling", "").strip(), int(f.get("duration_min") or 10), f["frequency"], f.get("machine_state", ""),
            int(f.get("people") or 1), f.get("sop_url", "").strip(), f.get("activity_type", "AM"), g.user["id"], now()))
        ex("UPDATE tags SET status='Contre-mesure', countermeasure=?, updated_at=?, resolved_at=COALESCE(resolved_at, ?) WHERE id=?",
           (f["countermeasure"].strip(), now(), now(), tag_id))
        ex("INSERT INTO tag_history(tag_id,user_id,from_status,to_status,comment,at) VALUES(?,?,?,?,?,?)",
           (tag_id, g.user["id"], t["status"], "Contre-mesure", "Contre-mesure enregistrée (à intégrer dans une checklist)", now()))
        flash("Contre-mesure enregistrée dans la liste des contre-mesures. Elle pourra être intégrée dans une checklist CILT.", "success")
        return redirect(url_for("tag_detail", tag_id=tag_id))
    v = {"organe": t["organe"], "frequency": "Hebdomadaire", "duration_min": 10, "people": 1,
         "activity_type": "RCM" if t["type"] == "RCM" else "AM", "machine_state": MACHINE_STATES[0], "cil_type": "I"}
    return render_template("countermeasure.html", t=t, v=v)


@app.route("/tags/<int:tag_id>/edit", methods=["GET", "POST"])
def tag_edit(tag_id):
    if not is_admin():
        abort(403)
    t = load_tag(tag_id)
    equipment = q("SELECT e.*, s.name sga_name FROM equipment e LEFT JOIN sga s ON s.id=e.sga_id ORDER BY e.position, e.name")
    if request.method == "POST":
        f = request.form
        e = q("SELECT * FROM equipment WHERE id=?", (f.get("equipment_id") or 0,), one=True)
        if f.get("type") not in TAG_TYPE or f.get("anomaly_type") not in ANOMALY_TYPES or not e \
                or not f.get("description", "").strip() or f.get("status") not in STATUSES:
            flash("Type, type d'anomalie, machine, description et statut sont obligatoires.", "error")
            return redirect(url_for("tag_edit", tag_id=tag_id))
        notif = clean_ref(f.get("sap_notif"))
        if notif and q("SELECT 1 FROM tags WHERE sap_notif=? AND id!=?", (notif, tag_id), one=True):
            flash("Ce N° d'avis est déjà rattaché à un autre TAG.", "error")
            return redirect(url_for("tag_edit", tag_id=tag_id))
        is_sap = f["type"] in sap.SAP_TYPES
        if is_sap and f["status"] != "Nouveau" and not notif:
            flash("Un TAG RCM / Safety hors « Découverte » doit avoir un N° d'avis SAP.", "error")
            return redirect(url_for("tag_edit", tag_id=tag_id))
        ot = clean_ref(f.get("ot_number")) if is_sap else ""
        photo = t["photo"]
        if f.get("remove_photo"):
            photo = ""
        new_photo = save_photo(request.files.get("photo"))
        if new_photo:
            photo = new_photo
        ex("""UPDATE tags SET type=?, anomaly_type=?, equipment_id=?, sga_id=?, line_id=?, organe=?, description=?, photo=?,
              observed_by=?, observed_at=?, treated_by=?, status=?, sap_notif=?, ot_number=?, sap_order=?, countermeasure=?,
              validated_by=CASE WHEN ?='' THEN NULL ELSE COALESCE(validated_by, ?) END,
              validated_at=CASE WHEN ?='' THEN NULL ELSE COALESCE(validated_at, ?) END, updated_at=? WHERE id=?""",
           (f["type"], f["anomaly_type"], e["id"], e["sga_id"], e["line_id"], f.get("organe", "").strip(), f["description"].strip(),
            photo, f.get("observed_by") or None, f.get("observed_at") or None, f.get("treated_by") or None, f["status"],
            notif or None, ot, ot or None, f.get("countermeasure", "").strip() or None,
            notif, g.user["id"], notif, now(), now(), tag_id))
        ex("INSERT INTO tag_history(tag_id,user_id,from_status,to_status,comment,at) VALUES(?,?,?,?,?,?)",
           (tag_id, g.user["id"], t["status"], f["status"], "TAG modifié par l'administrateur", now()))
        flash("%s modifié." % t["tag_code"], "success")
        return redirect(url_for("tag_detail", tag_id=tag_id))
    return render_template("tag_edit.html", t=t, equipment=equipment, users=active_users())


@app.route("/tags/<int:tag_id>/delete", methods=["POST"])
def tag_delete(tag_id):
    if not is_admin():
        abort(403)
    t = load_tag(tag_id)
    for sql in ("UPDATE task_items SET tag_id=NULL WHERE tag_id=?", "UPDATE pdr_requests SET tag_id=NULL WHERE tag_id=?",
                "UPDATE checklists SET origin_tag_id=NULL WHERE origin_tag_id=?", "UPDATE countermeasures SET tag_id=NULL WHERE tag_id=?",
                "UPDATE sap_log SET tag_id=NULL WHERE tag_id=?", "DELETE FROM tag_history WHERE tag_id=?", "DELETE FROM tags WHERE id=?"):
        ex(sql, (tag_id,))
    if t["photo"]:
        try:
            os.remove(os.path.join(UPLOADS, t["photo"]))
        except OSError:
            pass
        pg_store.delete_file(t["photo"])
    flash("%s supprimé." % t["tag_code"], "success")
    return redirect(url_for("tags_board"))


# ---------------------------------------------------------------------------
# Checklists CILT (gestion), distribution des tâches et contre-mesures
# ---------------------------------------------------------------------------
def add_months(d, n):
    y, m = divmod(d.month - 1 + n, 12)
    y += d.year
    m += 1
    last = [31, 29 if (y % 4 == 0 and (y % 100 != 0 or y % 400 == 0)) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31][m - 1]
    return d.replace(year=y, month=m, day=min(d.day, last))


def freq_add(d, freq):
    unit, n = FREQ_STEP.get(freq, ("days", 7))
    return d + dt.timedelta(days=n) if unit == "days" else add_months(d, n)


def occurrences(freq, start, end, last_due):
    """Dates d'échéance d'une tâche dans la période [start, end] (une tâche en retard est reprise au début de la période)."""
    anchor = start
    if last_due:
        try:
            anchor = freq_add(dt.date.fromisoformat(last_due[:10]), freq)
        except ValueError:
            anchor = start
    d = max(anchor, start)
    out = []
    while d <= end:
        out.append(d)
        d = freq_add(d, freq)
    return out


def plan_bounds(kind, ref):
    if kind == "mois":
        start = ref.replace(day=1)
        return start, add_months(start, 1) - dt.timedelta(days=1)
    if kind == "annee":
        return ref.replace(month=1, day=1), ref.replace(month=12, day=31)
    start = ref - dt.timedelta(days=ref.weekday())
    return start, start + dt.timedelta(days=6)


def checklist_rows(where="1=1", args=()):
    since = (today() - dt.timedelta(days=30)).isoformat()
    return q("""SELECT c.*, e.name eq_name, e.reference eq_ref, s.name sga_name, u.name user_name, u.role user_role,
               (SELECT COUNT(*) FROM checklist_items i WHERE i.checklist_id=c.id) n_items,
               (SELECT SUM(duration_min) FROM checklist_items i WHERE i.checklist_id=c.id) total_min,
               (SELECT GROUP_CONCAT(DISTINCT i.frequency) FROM checklist_items i WHERE i.checklist_id=c.id) freqs,
               (SELECT GROUP_CONCAT(DISTINCT i.cil_type) FROM checklist_items i WHERE i.checklist_id=c.id) types,
               (SELECT COUNT(*) FROM countermeasures m WHERE m.checklist_id=c.id) n_cm,
               (SELECT COUNT(*) FROM tasks k WHERE k.checklist_id=c.id AND k.due_date>=? AND k.due_date<=? AND k.status IN ('Réalisée','Confirmée')) n_done,
               (SELECT COUNT(*) FROM tasks k WHERE k.checklist_id=c.id AND k.due_date>=? AND k.due_date<=?) n_planned
               FROM checklists c JOIN equipment e ON e.id=c.equipment_id LEFT JOIN sga s ON s.id=e.sga_id
               LEFT JOIN users u ON u.id=c.assigned_to WHERE """ + where, [since, today().isoformat()] * 2 + list(args))


@app.route("/checklists")
def checklists():
    clause, args = in_clause("e.sga_id", scope_sga_ids())
    sga = request.args.get("sga", "")
    extra, ea = "", []
    if sga:
        extra, ea = " AND e.sga_id=?", [sga]
    rows = checklist_rows(clause + extra + " ORDER BY c.active DESC, s.name, e.position, c.code", args + ea)
    operators = q("SELECT * FROM users WHERE active=1 AND role='Opérateur' ORDER BY name")
    return render_template("checklists.html", rows=rows, sgas=scoped_sgas(), sel_sga=sga, operators=operators)


def item_rows_from_form(f):
    rows, errors = [], []
    ids, texts = f.getlist("item_id"), f.getlist("item_text")
    organes, types, tools = f.getlist("item_organe"), f.getlist("item_type"), f.getlist("item_tooling")
    durs, freqs = f.getlist("item_duration"), f.getlist("item_frequency")
    for i, txt in enumerate(texts):
        txt = (txt or "").strip()
        if not txt:
            continue
        typ = types[i] if i < len(types) else ""
        fq = freqs[i] if i < len(freqs) else ""
        try:
            dur = max(1, int(durs[i]))
        except (ValueError, IndexError):
            dur = None
        if typ not in CIL or fq not in FREQUENCIES or not dur:
            errors.append("Tâche « %s » : type, temps de réalisation et fréquence sont obligatoires." % txt[:40])
        rows.append({"id": ids[i] if i < len(ids) else "", "text": txt, "organe": (organes[i] if i < len(organes) else "").strip(),
                     "cil_type": typ, "tooling": (tools[i] if i < len(tools) else "").strip(), "duration_min": dur, "frequency": fq})
    if not rows:
        errors.append("Ajoutez au moins une tâche.")
    return rows, errors


def save_items(cid, rows):
    keep = []
    for pos, r in enumerate(rows, 1):
        vals = (pos, r["text"], r["organe"], r["cil_type"], r["tooling"], r["duration_min"], r["frequency"])
        if r["id"] and q("SELECT 1 FROM checklist_items WHERE id=? AND checklist_id=?", (r["id"], cid), one=True):
            ex("UPDATE checklist_items SET position=?, text=?, organe=?, cil_type=?, tooling=?, duration_min=?, frequency=? WHERE id=?",
               vals + (r["id"],))
            keep.append(int(r["id"]))
        else:
            keep.append(ex("INSERT INTO checklist_items(position,text,organe,cil_type,tooling,duration_min,frequency,checklist_id) "
                           "VALUES(?,?,?,?,?,?,?,?)", vals + (cid,)))
    clause, args = in_clause("id", keep)
    ex("DELETE FROM checklist_items WHERE checklist_id=? AND NOT (" + clause + ")", [cid] + args)


def checklist_form(cl=None):
    equipment = q("SELECT e.*, s.name sga_name FROM equipment e LEFT JOIN sga s ON s.id=e.sga_id ORDER BY e.position")
    users = q("SELECT u.*, (SELECT GROUP_CONCAT(s.name, ', ') FROM user_sga us JOIN sga s ON s.id=us.sga_id WHERE us.user_id=u.id) sgas "
              "FROM users u WHERE u.active=1 AND u.role IN ('Opérateur','Technicien','Leader SGA') ORDER BY u.role, u.name")
    if request.method == "POST":
        f = request.form
        rows, errors = item_rows_from_form(f)
        if not f.get("name") or not f.get("equipment_id"):
            errors.insert(0, "Le nom de la checklist et la machine sont obligatoires.")
        if errors:
            for er in errors:
                flash(er, "error")
            return render_template("checklist_form.html", cl=cl, v=f, items=rows, equipment=equipment, users=users)
        vals = (f["name"].strip(), f["equipment_id"], f.get("machine_state", ""), int(f.get("people") or 1), f.get("sop_url", "").strip(),
                f.get("activity_type", "AM"), f.get("assigned_to") or None, 1 if f.get("active", "1") == "1" else 0)
        if cl:
            ex("UPDATE checklists SET name=?, equipment_id=?, machine_state=?, people=?, sop_url=?, activity_type=?, assigned_to=?, active=? "
               "WHERE id=?", vals + (cl["id"],))
            cid = cl["id"]
        else:
            n = q("SELECT COUNT(*)+1 n FROM checklists WHERE code LIKE 'CILT-%'", one=True)["n"]
            cid = ex("INSERT INTO checklists(name,equipment_id,machine_state,people,sop_url,activity_type,assigned_to,active,code,created_at) "
                     "VALUES(?,?,?,?,?,?,?,?,?,?)", vals + ("CILT-%03d" % n, now()))
        save_items(cid, rows)
        cm_id = f.get("from_cm")
        if cm_id:
            ex("UPDATE countermeasures SET status='Intégrée', checklist_id=?, integrated_at=?, integrated_by=? WHERE id=?",
               (cid, now(), g.user["id"], cm_id))
        flash("Checklist enregistrée. Ses tâches seront distribuées depuis « Distribution des tâches ».", "success")
        return redirect(url_for("checklists"))
    v = dict(cl) if cl else {"people": 1, "activity_type": "AM", "active": 1, "equipment_id": request.args.get("equipment_id", "")}
    items = [dict(i) for i in q("SELECT * FROM checklist_items WHERE checklist_id=? ORDER BY position", (cl["id"],))] if cl else []
    return render_template("checklist_form.html", cl=cl, v=v, items=items, equipment=equipment, users=users)


@app.route("/checklists/new", methods=["GET", "POST"])
@require("checklists_manage")
def checklist_new():
    return checklist_form()


@app.route("/checklists/<int:cid>/edit", methods=["GET", "POST"])
@require("checklists_manage")
def checklist_edit(cid):
    cl = q("SELECT * FROM checklists WHERE id=?", (cid,), one=True) or abort(404)
    return checklist_form(cl)


@app.route("/checklists/<int:cid>/transfer", methods=["POST"])
@require("checklists_manage")
def checklist_transfer(cid):
    cl = q("SELECT * FROM checklists WHERE id=?", (cid,), one=True) or abort(404)
    assignee = request.form.get("assigned_to") or cl["assigned_to"]
    ex("UPDATE checklists SET transferred=1, transfer_date=?, activity_type='AM', assigned_to=? WHERE id=?",
       (today().isoformat(), assignee, cid))
    ex("UPDATE checklist_items SET last_assignee=? WHERE checklist_id=?", (assignee, cid))
    flash("Activité transférée de RCM vers AM.", "success")
    return redirect(url_for("checklists"))


@app.route("/checklists/<int:cid>/history")
def checklist_history(cid):
    cl = q("SELECT c.*, e.name eq_name FROM checklists c JOIN equipment e ON e.id=c.equipment_id WHERE c.id=?", (cid,), one=True) or abort(404)
    rows = task_rows("k.checklist_id=? ORDER BY k.due_date DESC", (cid,))
    items = q("SELECT * FROM checklist_items WHERE checklist_id=? ORDER BY position", (cid,))
    return render_template("checklist_history.html", cl=cl, rows=rows, items=items)


@app.route("/api/organes")
def api_organes():
    eid = request.args.get("equipment_id", 0)
    rows = q("""SELECT organe FROM (SELECT organe FROM tags WHERE equipment_id=? UNION SELECT i.organe FROM checklist_items i
               JOIN checklists c ON c.id=i.checklist_id WHERE c.equipment_id=? UNION SELECT organe FROM countermeasures WHERE equipment_id=?)
               WHERE organe IS NOT NULL AND organe!='' ORDER BY organe""", (eid, eid, eid))
    return jsonify([r["organe"] for r in rows])


# --- Distribution des tâches (administrateur) ------------------------------
def planning_rows(kind, start, end):
    freqs = PLAN_PERIODS[kind][1]
    clause, fargs = in_clause("i.frequency", freqs)
    items = q("""SELECT i.*, c.name cl_name, c.code cl_code, c.assigned_to cl_assignee, c.equipment_id, e.name eq_name, e.reference eq_ref,
                 s.name sga_name FROM checklist_items i JOIN checklists c ON c.id=i.checklist_id JOIN equipment e ON e.id=c.equipment_id
                 LEFT JOIN sga s ON s.id=e.sga_id WHERE c.active=1 AND """ + clause + " ORDER BY s.name, e.position, c.code, i.position", fargs)
    out = []
    for it in items:
        # une tâche quotidienne n'est pas créée pour les jours déjà passés de la période
        s0 = max(start, today()) if it["frequency"] == "Quotidienne" and start <= today() <= end else start
        dates = occurrences(it["frequency"], s0, end, it["last_due"])
        if dates:
            out.append({"i": it, "dates": dates, "assignee": it["last_assignee"] or it["cl_assignee"]})
    return out


@app.route("/planning", methods=["GET", "POST"])
@require("checklists_manage")
def planning():
    kind = request.values.get("period", "semaine")
    if kind not in PLAN_PERIODS:
        kind = "semaine"
    try:
        ref = dt.date.fromisoformat(request.values.get("date", ""))
    except ValueError:
        ref = today()
    start, end = plan_bounds(kind, ref)
    if request.method == "POST":
        selected = set(request.form.getlist("sel"))
        groups = {}
        missing = 0
        for r in planning_rows(kind, start, end):
            it = r["i"]
            if str(it["id"]) not in selected:
                continue
            who = request.form.get("assignee_%d" % it["id"])
            if not who:
                missing += 1
                continue
            for d in r["dates"]:
                groups.setdefault((it["checklist_id"], int(who), d), []).append(it)
            ex("UPDATE checklist_items SET last_due=?, last_assignee=? WHERE id=?", (r["dates"][-1].isoformat(), who, it["id"]))
        for (cid, who, d), its in sorted(groups.items(), key=lambda kv: (kv[0][2], kv[0][0])):
            tid = ex("INSERT INTO tasks(checklist_id,assigned_to,due_date,status) VALUES(?,?,?, 'À réaliser')", (cid, who, d.isoformat()))
            for pos, it in enumerate(its, 1):
                ex("""INSERT INTO task_items(task_id,position,text,checklist_item_id,organe,cil_type,tooling,duration_min,frequency)
                      VALUES(?,?,?,?,?,?,?,?,?)""", (tid, pos, it["text"], it["id"], it["organe"], it["cil_type"], it["tooling"],
                                                     it["duration_min"], it["frequency"]))
        if missing:
            flash("%d tâche(s) non distribuée(s) : choisissez la personne concernée." % missing, "error")
        flash("%d ordre(s) de travail distribué(s) pour la période du %s au %s." % (len(groups), start.strftime("%d/%m/%Y"),
                                                                                  end.strftime("%d/%m/%Y")), "success")
        return redirect(url_for("planning", period=kind, date=ref.isoformat()))
    rows = planning_rows(kind, start, end)
    people = q("""SELECT u.*, (SELECT GROUP_CONCAT(s.name, ', ') FROM user_sga us JOIN sga s ON s.id=us.sga_id WHERE us.user_id=u.id) sgas
                  FROM users u WHERE u.active=1 AND u.role IN ('Opérateur','Technicien','Leader SGA','Chef d''équipe Production',
                  'Chef d''équipe Maintenance') ORDER BY u.role, u.name""")
    done = task_rows("k.due_date>=? AND k.due_date<=? ORDER BY k.due_date, u.name", (start.isoformat(), end.isoformat()))
    prev_ref = start - dt.timedelta(days=1)
    next_ref = end + dt.timedelta(days=1)
    return render_template("planning.html", kind=kind, rows=rows, people=people, done=done, start=start, end=end,
                           ref=ref, prev_ref=prev_ref, next_ref=next_ref, periods=PLAN_PERIODS)


@app.route("/planning/task/<int:task_id>/cancel", methods=["POST"])
@require("checklists_manage")
def planning_cancel(task_id):
    t = q("SELECT * FROM tasks WHERE id=?", (task_id,), one=True) or abort(404)
    started = q("SELECT COUNT(*) n FROM task_items WHERE task_id=? AND ok IS NOT NULL", (task_id,), one=True)["n"]
    if t["status"] != "À réaliser" or started:
        flash("Cet ordre de travail est déjà commencé : il ne peut plus être annulé.", "error")
    else:
        item_ids = [r["checklist_item_id"] for r in q("SELECT checklist_item_id FROM task_items WHERE task_id=?", (task_id,)) if r["checklist_item_id"]]
        ex("DELETE FROM tasks WHERE id=?", (task_id,))
        for iid in item_ids:
            last = q("SELECT MAX(k.due_date) d FROM task_items ti JOIN tasks k ON k.id=ti.task_id WHERE ti.checklist_item_id=?", (iid,), one=True)
            ex("UPDATE checklist_items SET last_due=? WHERE id=?", (last["d"] if last else None, iid))
        flash("Ordre de travail annulé : ses tâches redeviennent à distribuer.", "success")
    return redirect(request.referrer or url_for("planning"))


# --- Contre-mesures --------------------------------------------------------
@app.route("/countermeasures")
def countermeasures():
    st = request.args.get("status", "À intégrer")
    clause, args = in_clause("e.sga_id", scope_sga_ids())
    extra, ea = "", []
    if st in ("À intégrer", "Intégrée"):
        extra, ea = " AND m.status=?", [st]
    rows = q("""SELECT m.*, t.tag_code, t.type tag_type, e.name eq_name, e.reference eq_ref, s.name sga_name, c.code cl_code, c.name cl_name,
               u.name author FROM countermeasures m JOIN equipment e ON e.id=m.equipment_id LEFT JOIN sga s ON s.id=e.sga_id
               LEFT JOIN tags t ON t.id=m.tag_id LEFT JOIN checklists c ON c.id=m.checklist_id LEFT JOIN users u ON u.id=m.created_by
               WHERE """ + ("1=1" if is_admin() else clause) + extra + " ORDER BY m.created_at DESC", ([] if is_admin() else args) + ea)
    counts = {r["status"]: r["n"] for r in q("SELECT status, COUNT(*) n FROM countermeasures GROUP BY status")}
    return render_template("countermeasures.html", rows=rows, st=st, counts=counts)


@app.route("/countermeasures/<int:mid>/integrate", methods=["GET", "POST"])
@require("checklists_manage")
def countermeasure_integrate(mid):
    m = q("""SELECT m.*, t.tag_code, e.name eq_name, e.reference eq_ref FROM countermeasures m JOIN equipment e ON e.id=m.equipment_id
             LEFT JOIN tags t ON t.id=m.tag_id WHERE m.id=?""", (mid,), one=True) or abort(404)
    cls = q("SELECT * FROM checklists WHERE equipment_id=? AND active=1 ORDER BY code", (m["equipment_id"],))
    target = request.values.get("checklist_id", "")
    tasks = [x for x in (m["tasks"] or m["countermeasure"]).splitlines() if x.strip()]
    per = max(1, round((m["duration_min"] or 10) / max(1, len(tasks))))
    new_items = [{"id": "", "text": x.strip(), "organe": m["organe"] or "", "cil_type": m["cil_type"] or "I", "tooling": m["tooling"] or "",
                  "duration_min": per, "frequency": m["frequency"] or "Hebdomadaire"} for x in tasks]
    if request.method == "GET" and target == "":
        return render_template("countermeasure_integrate.html", m=m, cls=cls)
    if request.method == "POST":
        cl = None if target == "new" else (q("SELECT * FROM checklists WHERE id=? AND equipment_id=?", (target, m["equipment_id"]), one=True) or abort(404))
        return checklist_form(cl)
    if target == "new":
        v = {"name": "CILT %s" % (m["eq_name"] or ""), "equipment_id": m["equipment_id"], "machine_state": m["machine_state"],
             "people": m["people"] or 1, "sop_url": m["sop_url"], "activity_type": m["activity_type"] or "AM", "active": 1, "from_cm": mid}
        return render_template("checklist_form.html", cl=None, v=v, items=new_items, cm=m,
                               equipment=q("SELECT e.*, s.name sga_name FROM equipment e LEFT JOIN sga s ON s.id=e.sga_id ORDER BY e.position"),
                               users=q("SELECT u.*, '' sgas FROM users u WHERE u.active=1 AND u.role IN ('Opérateur','Technicien','Leader SGA') ORDER BY u.role, u.name"))
    cl = q("SELECT * FROM checklists WHERE id=? AND equipment_id=?", (target, m["equipment_id"]), one=True) or abort(404)
    items = [dict(i) for i in q("SELECT * FROM checklist_items WHERE checklist_id=? ORDER BY position", (cl["id"],))] + new_items
    v = dict(cl)
    v["from_cm"] = mid
    return render_template("checklist_form.html", cl=cl, v=v, items=items, cm=m,
                           equipment=q("SELECT e.*, s.name sga_name FROM equipment e LEFT JOIN sga s ON s.id=e.sga_id ORDER BY e.position"),
                           users=q("SELECT u.*, '' sgas FROM users u WHERE u.active=1 AND u.role IN ('Opérateur','Technicien','Leader SGA') ORDER BY u.role, u.name"))


# ---------------------------------------------------------------------------
# Documentation
# ---------------------------------------------------------------------------
@app.route("/documentation")
def documentation():
    eq_id = request.args.get("equipment", "")
    plants = q("SELECT * FROM plants ORDER BY name")
    tree = []
    for p in plants:
        lines = []
        for l in q("SELECT * FROM lines WHERE plant_id=? ORDER BY name", (p["id"],)):
            sgas = []
            for s in q("SELECT * FROM sga WHERE line_id=? ORDER BY name", (l["id"],)):
                sgas.append({"s": s, "eq": q("SELECT * FROM equipment WHERE sga_id=? ORDER BY position", (s["id"],))})
            orphans = q("SELECT * FROM equipment WHERE line_id=? AND sga_id IS NULL ORDER BY position", (l["id"],))
            lines.append({"l": l, "sgas": sgas, "orphans": orphans})
        tree.append({"p": p, "lines": lines})
    eq = None
    docs = []
    if eq_id:
        eq = q("SELECT e.*, s.name sga_name, l.name line_name, p.name plant_name FROM equipment e LEFT JOIN sga s ON s.id=e.sga_id "
               "JOIN lines l ON l.id=e.line_id JOIN plants p ON p.id=l.plant_id WHERE e.id=?", (eq_id,), one=True)
        docs = q("SELECT * FROM documents WHERE equipment_id=? ORDER BY doc_type, title", (eq_id,))
        cls = q("SELECT * FROM checklists WHERE equipment_id=? AND active=1 ORDER BY code", (eq_id,))
    else:
        cls = []
    dtype = request.args.get("type", "")
    return render_template("documentation.html", tree=tree, eq=eq, docs=docs, cls=cls, dtype=dtype,
                           all_equipment=q("SELECT * FROM equipment ORDER BY position"))


@app.route("/documentation/add", methods=["POST"])
@require("docs_manage")
def doc_add():
    f = request.form
    if not f.get("equipment_id") or not f.get("title"):
        flash("Équipement et titre obligatoires.", "error")
    else:
        ex("INSERT INTO documents(equipment_id,title,doc_type,version,owner,url,updated_at) VALUES(?,?,?,?,?,?,?)",
           (f["equipment_id"], f["title"], f.get("doc_t
